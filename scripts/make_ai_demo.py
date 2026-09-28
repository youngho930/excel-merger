"""AI 매칭 체험용 샘플을 만든다: samples/ai_demo/

기존 샘플(scripts/make_samples.py, 정답지 37건)과 별개다. 그 파일들은 건드리지 않는다.

- 수입검사 시나리오용 파일 1개. 열 이름 일부가 동의어 사전에 없다
  ("자재 식별번호", "공급사", "입하 수량", "불량 개수", "합부") -> 동의어 매칭으로는 못 찾고 AI 추천이 필요하다.
- demo.yaml: 이 폴더가 어느 시나리오의 체험 파일인지 적은 설명서. 앱은 이 파일이 있는 폴더를
  시나리오 샘플 폴더로 보지 않는다.
- 심은 오류 2건 (AI 추천을 그대로 받아들였을 때):
  4행 입하 수량 "20박스"(형식 오류), 7행 합부 "보류"(허용값 아닌 값)

실행: python scripts/make_ai_demo.py
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "samples" / "ai_demo"
FILE_NAME = "수입검사_AI체험_신규협력사.xlsx"

HEADERS = ["자재 식별번호", "자재명", "LOT No", "공급사", "검사일자", "입하 수량", "검사 수량", "불량 개수", "합부"]
ROWS = [
    ("NP-3001", "브래킷 A", "NP260915-01", "새한테크", datetime(2026, 9, 15), 500, 20, 0, "합격"),
    ("NP-3002", "브래킷 B", "NP260915-02", "새한테크", datetime(2026, 9, 15), 300, 13, 1, "합격"),
    ("NP-3003", "커버 플레이트", "NP260916-01", "새한테크", datetime(2026, 9, 16), "20박스", 8, 0, "합격"),
    ("NP-3004", "고정 핀", "NP260916-02", "새한테크", datetime(2026, 9, 16), 1000, 32, 3, "불합격"),
    ("NP-3005", "스페이서", "NP260917-01", "새한테크", datetime(2026, 9, 17), 800, 32, 0, "합격"),
    ("NP-3006", "가이드 레일", "NP260917-02", "새한테크", datetime(2026, 9, 17), 120, 8, 0, "보류"),
    ("NP-3007", "캡", "NP260918-01", "새한테크", datetime(2026, 9, 18), 2000, 50, 2, "합격"),
]
EXPECTED = [
    {"row": 4, "column": "입하 수량", "standard": "입고수량", "kind": "형식 오류"},
    {"row": 7, "column": "합부", "standard": "판정", "kind": "허용값 아닌 값"},
]
# AI가 찾아야 할 짝 (원본 열 -> 기준열). 테스트의 가짜 응답과 설명에 쓴다.
AI_TARGETS = {"자재 식별번호": "품목코드", "공급사": "협력사", "입하 수량": "입고수량",
              "불량 개수": "불량수량", "합부": "판정"}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "수입검사"
    ws.append(HEADERS)
    for r in ROWS:
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    wb.save(OUT / FILE_NAME)

    manifest = {
        "scenario": "incoming_inspection",
        "description": "동의어 사전에 없는 열 이름을 쓴 신규 협력사의 수입검사 결과표",
        "files": [FILE_NAME],
        "ai_targets": AI_TARGETS,
        "expected_errors": EXPECTED,
    }
    (OUT / "demo.yaml").write_text(
        "# scripts/make_ai_demo.py 가 만든다. 직접 고치지 말고 스크립트를 고친 뒤 다시 실행하세요.\n"
        + yaml.safe_dump(manifest, allow_unicode=True, sort_keys=False), encoding="utf-8")
    print(f"만듦: {OUT / FILE_NAME}")


if __name__ == "__main__":
    main()
