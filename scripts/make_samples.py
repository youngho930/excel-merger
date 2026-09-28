"""샘플 엑셀 9개와 정답지(samples/expected_errors.md, samples/expected_errors.json)를 만든다.

실행: python scripts/make_samples.py

같은 시나리오 안에서도 파일마다 열 이름·열 순서·날짜 형식을 다르게 하고,
오류를 일부러 심는다. 심은 오류는 그대로 정답지에 기록되므로
(사람이 읽는 .md, 테스트가 자동 대조하는 .json)
샘플과 정답지는 항상 이 스크립트로 함께 다시 만든다.
"""

import copy
import json
import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"

BLANK = None  # 빈칸 오류를 나타내는 값

FICTION_NOTE = "샘플의 회사명·인명은 모두 가상입니다. 실제 회사나 인물과 관계없습니다."

# 오류 종류 (CLAUDE.md "검증 종류"와 같은 이름)
REQUIRED = "필수값 빈칸"
FORMAT = "형식 오류"
RANGE = "범위 밖 값"
ALLOWED = "허용값 아닌 값"
DUPLICATE = "중복 행"


# ---------------------------------------------------------------- 날짜 표기
def as_datetime(d):
    return datetime(d.year, d.month, d.day)


def dotted(d):
    return f"{d:%Y.%m.%d}"


def slashed(d):
    return f"{d:%Y/%m/%d}"


def dotted_short(d):
    return f"{d.year}.{d.month}.{d.day}"


def korean(d):
    return f"{d.year}년 {d.month}월 {d.day}일"


def month_dash(d):
    return f"{d:%Y-%m}"


def month_korean(d):
    return f"{d.year}년 {d.month}월"


# ---------------------------------------------------------------- 파일 스펙
class SampleFile:
    """한 샘플 파일: 열 이름·순서·날짜 표기·머리글 위치와 데이터, 심은 오류."""

    def __init__(self, filename, sheet, headers, date_cols, date_fmt, rows,
                 title_lines=None, date_number_format="yyyy-mm-dd"):
        self.filename = filename
        self.sheet = sheet
        self.headers = headers          # [(기준명, 파일에 쓰인 열 이름)] 순서대로
        self.date_cols = date_cols      # 날짜로 표기할 기준명 목록
        self.date_fmt = date_fmt        # date -> 셀 값
        self.date_number_format = date_number_format
        self.clean = rows               # 오류를 심기 전 원본 (기준명 -> 값)
        self.rows = copy.deepcopy(rows)
        self.title_lines = title_lines or []
        self.errors = []                # 정답지에 들어갈 기록
        self.touched = set()            # 오류를 심은 행 번호(0부터)

    @property
    def header_row(self):
        # 머리글 위 줄이 있으면 빈 줄 없이 바로 이어서 머리글이 온다
        return len(self.title_lines) + 1

    def excel_row(self, idx):
        return self.header_row + 1 + idx

    def col_name(self, key):
        return dict(self.headers)[key]

    def plant(self, idx, key, value, kind, desc, shown=None):
        self.rows[idx][key] = value
        self.touched.add(idx)
        self.errors.append({
            "row": self.excel_row(idx),
            "col": self.col_name(key),
            "key": key,
            "kind": kind,
            "value": shown if shown is not None else value,
            "raw": value,
            "desc": desc,
            "duplicate_of": None,
        })

    def plant_duplicate(self, idx, source, src_idx, dup_keys):
        """source 파일의 src_idx 행(오류 없는 행)을 이 파일 idx 자리에 그대로 넣는다."""
        assert src_idx not in source.touched, "오류를 심은 행은 중복 원본으로 쓰지 않는다"
        self.rows[idx] = copy.deepcopy(source.clean[src_idx])
        self.clean[idx] = copy.deepcopy(source.clean[src_idx])
        self.touched.add(idx)
        key_text = ", ".join(f"{k}={fmt_key(self.rows[idx][k])}" for k in dup_keys)
        self.errors.append({
            "row": self.excel_row(idx),
            "col": "(행 전체)",
            "key": " + ".join(dup_keys),
            "kind": DUPLICATE,
            "value": key_text,
            "raw": None,
            "desc": f"{source.filename} {source.excel_row(src_idx)}행과 중복기준 값이 같음",
            "duplicate_of": {"file": source.filename, "row": source.excel_row(src_idx)},
        })

    def save(self, folder):
        wb = Workbook()
        ws = wb.active
        ws.title = self.sheet
        # 서식은 모든 행을 append한 뒤에 준다. 먼저 셀을 건드리면 빈 1행이 생겨
        # 머리글이 한 줄 밀리고 정답지 행 번호와 어긋난다.
        for line in self.title_lines:
            ws.append([line])
        ws.append([name for _, name in self.headers])
        for row in self.rows:
            values = []
            for key, _ in self.headers:
                v = row.get(key)
                if key in self.date_cols and isinstance(v, date):
                    v = self.date_fmt(v)
                values.append(v)
            ws.append(values)
        if self.title_lines:
            ws.cell(row=1, column=1).font = Font(bold=True, size=14)
        for cell in ws[self.header_row]:
            cell.font = Font(bold=True)
        for col_cells in ws.iter_cols(min_row=self.header_row):
            for cell in col_cells:
                if isinstance(cell.value, datetime):
                    cell.number_format = self.date_number_format
        for i, _ in enumerate(self.headers, start=1):
            ws.column_dimensions[ws.cell(row=self.header_row, column=i).column_letter].width = 14
        wb.save(folder / self.filename)


def fmt_key(v):
    if isinstance(v, date):
        return v.isoformat()
    return str(v)


# ---------------------------------------------------------------- 1. 수입검사
PARTS = [
    ("HB-1001", "브라켓 A"), ("HB-1002", "브라켓 B"), ("HB-2010", "샤프트 20mm"),
    ("HB-2011", "샤프트 25mm"), ("HB-3100", "하우징 커버"), ("HB-3101", "하우징 베이스"),
    ("HB-4005", "가스켓"), ("HB-4006", "O링 30"), ("HB-5200", "볼트 M6"),
    ("HB-5201", "너트 M6"), ("HB-6300", "스프링 와셔"), ("HB-7010", "커넥터 4P"),
]


def inspection_rows(rng, supplier, lot_prefix, count, start):
    rows = []
    for i in range(count):
        code, name = rng.choice(PARTS)
        d = start + timedelta(days=i // 3)
        received = rng.choice([100, 200, 300, 500, 1000])
        inspected = min(received, rng.choice([13, 20, 32, 50]))
        defects = rng.choice([0, 0, 0, 0, 1, 2])
        rows.append({
            "품목코드": code, "품목명": name,
            "로트번호": f"{lot_prefix}{d:%y%m%d}-{i + 1:02d}",
            "협력사": supplier, "검사일": d,
            "입고수량": received, "검사수량": inspected, "불량수량": defects,
            "판정": "합격" if defects == 0 else "불합격",
        })
    return rows


def build_incoming(rng):
    dup_keys = ["품목코드", "로트번호"]

    a = SampleFile(
        "수입검사_대성정공.xlsx", "Sheet1",
        [("품목코드", "품번"), ("품목명", "품명"), ("로트번호", "로트번호"), ("협력사", "협력사"),
         ("검사일", "검사일"), ("입고수량", "입고수량"), ("검사수량", "검사수량"),
         ("불량수량", "불량수량"), ("판정", "판정")],
        ["검사일"], as_datetime,
        inspection_rows(rng, "대성정공", "DS", 15, date(2026, 9, 1)),
    )
    a.plant(4, "로트번호", BLANK, REQUIRED, "필수 열 '로트번호'가 비어 있음")
    a.plant(7, "입고수량", "12개", FORMAT, "정수 칸에 문자('개')가 섞임")
    a.plant(10, "불량수량", -2, RANGE, "수량이 음수 (허용 범위 0~100000)")
    a.plant(12, "판정", "합격 ", ALLOWED, "'합격' 뒤에 공백이 있어 허용값과 다름", shown='"합격 " (뒤에 공백)')

    b = SampleFile(
        "수입검사_미래부품.xlsx", "검사결과",
        [("검사일", "검사일자"), ("협력사", "업체명"), ("품목코드", "품목코드"), ("품목명", "품목명"),
         ("로트번호", "LOT No"), ("판정", "검사결과"), ("입고수량", "입고량"),
         ("검사수량", "샘플수량"), ("불량수량", "불량수")],
        ["검사일"], dotted,
        inspection_rows(rng, "(주)미래부품", "MR", 12, date(2026, 9, 2)),
    )
    b.plant(2, "품목코드", BLANK, REQUIRED, "필수 열 '품목코드'이 비어 있음")
    b.plant(5, "판정", "OK", ALLOWED, "허용값(합격/불합격)이 아닌 'OK'")
    b.plant(6, "검사일", "2026.13.02", FORMAT, "존재하지 않는 날짜(13월)")
    b.plant(9, "검사수량", "20EA", FORMAT, "정수 칸에 문자('EA')가 섞임")
    b.plant_duplicate(11, a, 1, dup_keys)

    c = SampleFile(
        "수입검사_한일정밀.xlsx", "Inspection",
        [("로트번호", "Lot"), ("품목코드", "Part No"), ("품목명", "Part Name"), ("협력사", "Supplier"),
         ("입고수량", "Recv Qty"), ("검사수량", "Insp Qty"), ("불량수량", "NG Qty"),
         ("판정", "Result"), ("검사일", "Insp. Date")],
        ["검사일"], slashed,
        inspection_rows(rng, "한일정밀", "HI", 18, date(2026, 9, 3)),
        title_lines=["한일정밀(주) 수입검사 성적서 (2026년 9월)", "작성일: 2026-09-10   작성자: 품질보증팀 박OO"],
    )
    c.plant(3, "입고수량", -50, RANGE, "수량이 음수 (허용 범위 0~100000)")
    c.plant(8, "판정", "Pass", ALLOWED, "허용값(합격/불합격)이 아닌 'Pass'")
    c.plant(11, "판정", BLANK, REQUIRED, "필수 열 '판정'이 비어 있음")
    c.plant(14, "검사수량", "32개", FORMAT, "정수 칸에 문자('개')가 섞임")
    c.plant_duplicate(16, b, 4, dup_keys)

    return [a, b, c]


# ---------------------------------------------------------------- 2. 재고실사
ITEMS = [
    ("RM-1001", "냉연강판 1.2T", "KG"), ("RM-1002", "냉연강판 1.6T", "KG"),
    ("RM-2001", "알루미늄 봉 20Φ", "M"), ("RM-2002", "알루미늄 봉 30Φ", "M"),
    ("PK-3001", "포장박스 대", "BOX"), ("PK-3002", "포장박스 소", "BOX"),
    ("PK-3003", "스트레치 필름", "ROLL"), ("CP-4001", "볼트 M8", "EA"),
    ("CP-4002", "너트 M8", "EA"), ("CP-4003", "평와셔 8", "EA"),
    ("CP-4004", "스프링 와셔 8", "EA"), ("CP-4005", "리벳 4mm", "EA"),
    ("AS-5001", "모터 조립체", "SET"), ("AS-5002", "센서 모듈", "SET"),
    ("AS-5003", "제어 보드", "EA"), ("AS-5004", "케이블 하네스", "EA"),
    ("CH-6001", "절삭유", "KG"), ("CH-6002", "방청유", "KG"),
]


def stock_rows(rng, warehouse, counters, count, day):
    rows = []
    for code, name, unit in rng.sample(ITEMS, count):
        book = rng.randint(10, 500)
        counted = book + rng.choice([0, 0, 0, 0, -1, 1, -2, 3])
        rows.append({
            "창고": warehouse, "품목코드": code, "품목명": name, "단위": unit,
            "전산수량": book, "실사수량": counted,
            "실사자": rng.choice(counters), "실사일": day,
        })
    return rows


def build_stock(rng):
    dup_keys = ["창고", "품목코드"]

    a = SampleFile(
        "재고실사_평택1창고.xlsx", "Sheet1",
        [("창고", "창고"), ("품목코드", "품번"), ("품목명", "품명"), ("단위", "단위"),
         ("전산수량", "전산수량"), ("실사수량", "실사수량"), ("실사자", "실사자"), ("실사일", "실사일")],
        ["실사일"], as_datetime,
        stock_rows(rng, "평택1창고", ["김철수", "이영희"], 15, date(2026, 9, 26)),
    )
    a.plant(2, "실사수량", BLANK, REQUIRED, "필수 열 '실사수량'이 비어 있음")
    a.plant(5, "실사수량", "12개", FORMAT, "정수 칸에 문자('개')가 섞임")
    a.plant(9, "전산수량", -5, RANGE, "수량이 음수 (허용 범위 0~1000000)")
    a.plant(12, "단위", "개", ALLOWED, "허용 단위(EA/BOX/KG/M/SET/ROLL)가 아닌 '개'")

    b = SampleFile(
        "재고실사_구미창고.xlsx", "재고",
        [("품목코드", "Item Code"), ("품목명", "Description"), ("창고", "Warehouse"), ("실사일", "Count Date"),
         ("단위", "UOM"), ("실사수량", "Count Qty"), ("전산수량", "Book Qty"), ("실사자", "Counter")],
        ["실사일"], dotted_short,
        stock_rows(rng, "구미창고", ["박민수", "최지은"], 12, date(2026, 9, 26)),
    )
    b.plant(1, "품목코드", BLANK, REQUIRED, "필수 열 '품목코드'이 비어 있음")
    b.plant(4, "실사일", "2026.9.31", FORMAT, "존재하지 않는 날짜(9월 31일)")
    b.plant(7, "실사수량", -3, RANGE, "수량이 음수 (허용 범위 0~1000000)")
    b.plant(10, "단위", "박스", ALLOWED, "허용 단위(EA/BOX/KG/M/SET/ROLL)가 아닌 '박스'")
    b.plant_duplicate(11, a, 3, dup_keys)

    c = SampleFile(
        "재고실사_부산물류센터.xlsx", "실사표",
        [("실사일", "실사일자"), ("창고", "창고명"), ("품목코드", "품목코드"), ("품목명", "자재명"),
         ("단위", "단위"), ("전산수량", "장부수량"), ("실사수량", "실재고"), ("실사자", "담당자")],
        ["실사일"], korean,
        stock_rows(rng, "부산물류센터", ["윤서진", "오민재"], 16, date(2026, 9, 27)),
        title_lines=["(주)한빛정밀 부산물류센터 재고실사표", "실사기준일: 2026-09-27   승인: 물류팀장"],
    )
    c.plant(3, "실사자", BLANK, REQUIRED, "필수 열 '실사자'가 비어 있음")
    c.plant(6, "전산수량", "약 120", FORMAT, "정수 칸에 문자('약')가 섞임")
    c.plant(10, "실사수량", -10, RANGE, "수량이 음수 (허용 범위 0~1000000)")
    c.plant_duplicate(14, b, 5, dup_keys)

    return [a, b, c]


# ---------------------------------------------------------------- 3. 월간실적
KPIS = {
    "영업팀": [("매출액(백만원)", 1200), ("신규고객수", 8), ("수주건수", 45), ("견적건수", 120),
             ("고객방문수", 60), ("매출채권회수율(%)", 95), ("수주잔고(백만원)", 3000),
             ("고객클레임건수", 2), ("신제품매출(백만원)", 150), ("해외매출(백만원)", 400),
             ("영업이익률(%)", 8)],
    "생산팀": [("생산량(EA)", 52000), ("설비가동률(%)", 85), ("설비고장시간(h)", 20),
             ("인시당생산성(EA)", 35), ("납기준수율(%)", 98), ("재공재고(EA)", 4000),
             ("잔업시간(h)", 300), ("공정불량률(ppm)", 500), ("에너지사용량(MWh)", 180),
             ("원가절감(백만원)", 30), ("안전사고건수", 0), ("작업자수", 64)],
    "품질팀": [("고객불량률(ppm)", 50), ("수입검사건수", 400), ("수입검사불합격건수", 6),
             ("시정조치완료율(%)", 90), ("내부심사건수", 4), ("계측기교정완료율(%)", 100),
             ("품질비용(백만원)", 45), ("협력사평가건수", 10), ("신뢰성시험건수", 12),
             ("고객감사대응건수", 1)],
}


def monthly_rows(rng, dept, month):
    rows = []
    for item, target in KPIS[dept]:
        actual = round(target * rng.uniform(0.85, 1.12), 1) if target else 0
        note = rng.choice(["", "", "", "전월 대비 개선", "목표 재검토 필요"])
        rows.append({"부서": dept, "월": month, "항목": item, "목표": target,
                     "실적": actual, "비고": note or None})
    return rows


def build_monthly(rng):
    dup_keys = ["부서", "월", "항목"]
    month = date(2026, 8, 1)

    a = SampleFile(
        "월간실적_영업팀.xlsx", "Sheet1",
        [("부서", "부서"), ("월", "월"), ("항목", "항목"), ("목표", "목표"), ("실적", "실적"), ("비고", "비고")],
        ["월"], as_datetime, monthly_rows(rng, "영업팀", month),
        date_number_format="yyyy-mm",
    )
    a.plant(2, "실적", BLANK, REQUIRED, "필수 열 '실적'이 비어 있음")
    a.plant(5, "목표", "95%", FORMAT, "실수 칸에 문자('%')가 섞임")
    a.plant(8, "부서", "영업 팀", ALLOWED, "허용 부서명이 아님 ('영업팀'에 공백)")

    b = SampleFile(
        "월간실적_생산팀.xlsx", "8월실적",
        [("월", "기준월"), ("항목", "KPI"), ("부서", "팀명"), ("실적", "Actual"), ("목표", "Target"), ("비고", "Remark")],
        ["월"], month_dash, monthly_rows(rng, "생산팀", month),
        title_lines=["2026년 8월 부서별 월간 실적 보고", "제출부서: 생산팀   제출일: 2026-09-03"],
    )
    b.plant(1, "항목", BLANK, REQUIRED, "필수 열 '항목'이 비어 있음")
    b.plant(4, "실적", -120, RANGE, "실적이 음수 (허용 범위 0 이상)")
    b.plant(7, "목표", "약 500", FORMAT, "실수 칸에 문자('약')가 섞임")

    c = SampleFile(
        "월간실적_품질팀.xlsx", "Sheet1",
        [("항목", "관리항목"), ("부서", "Dept"), ("월", "년월"), ("실적", "실적치"), ("목표", "목표치"), ("비고", "특이사항")],
        ["월"], month_korean, monthly_rows(rng, "품질팀", month) + [{}],
    )
    c.plant(3, "실적", "3건", FORMAT, "실수 칸에 문자('건')가 섞임")
    c.plant(6, "부서", "QA팀", ALLOWED, "허용 부서명이 아닌 'QA팀'")
    c.plant(9, "목표", -1, RANGE, "목표가 음수 (허용 범위 0 이상)")
    c.plant_duplicate(10, b, 8, dup_keys)

    return [a, b, c]


# ---------------------------------------------------------------- 정답지
SCENARIOS = [
    ("incoming_inspection", "협력사 수입검사 결과", build_incoming),
    ("stock_count", "지점·창고 재고실사", build_stock),
    ("monthly_report", "부서별 월간 실적", build_monthly),
]

DATE_NOTES = {
    as_datetime: "엑셀 날짜 셀",
    dotted: "문자 `YYYY.MM.DD`",
    slashed: "문자 `YYYY/MM/DD`",
    dotted_short: "문자 `YYYY.M.D`",
    korean: "문자 `YYYY년 M월 D일`",
    month_dash: "문자 `YYYY-MM`",
    month_korean: "문자 `YYYY년 M월`",
}


def md_value(v):
    if v is None:
        return "(빈칸)"
    return str(v).replace("|", "\\|")


def write_answer_key(results):
    total = sum(len(f.errors) for _, _, files in results for f in files)
    lines = [
        "# 정답지: 샘플에 심은 오류 목록",
        "",
        "> 이 파일은 `scripts/make_samples.py`가 샘플과 함께 자동으로 만듭니다. 직접 고치지 말고 스크립트를 고친 뒤 다시 실행하세요.",
        "",
        f"> {FICTION_NOTE}",
        "",
        "- **행 번호**는 엑셀 화면에 보이는 실제 행 번호입니다 (머리글 행이 3행인 파일은 데이터가 4행부터 시작).",
        "- **열 이름**은 그 파일에 실제로 쓰인 이름이고, 괄호 안은 시나리오의 기준명입니다.",
        "- **중복 행**은 뒤에 나온 행을 오류로 적고, 설명에 먼저 나온 행을 적었습니다. 엔진이 두 행 모두 표시해도 정답으로 봅니다.",
        "- 파일마다 다른 날짜 표기(아래 표)와 열 이름·열 순서 차이는 **오류가 아닙니다.** 엔진이 알아서 맞춰야 합니다.",
        "",
        f"총 **{total}건**",
        "",
        "## 파일별 요약",
        "",
        "| 시나리오 | 파일 | 머리글 행 | 데이터 행 수 | 날짜 표기 | 심은 오류 수 |",
        "|---|---|---|---|---|---|",
    ]
    for folder, title, files in results:
        for f in files:
            lines.append(
                f"| {title} | `{folder}/{f.filename}` | {f.header_row}행 | {len(f.rows)} "
                f"| {DATE_NOTES[f.date_fmt]} | {len(f.errors)} |"
            )

    kinds = [REQUIRED, FORMAT, RANGE, ALLOWED, DUPLICATE]
    lines += ["", "## 오류 종류별 개수", "", "| 오류 종류 | 개수 |", "|---|---|"]
    for k in kinds:
        n = sum(1 for _, _, files in results for f in files for e in f.errors if e["kind"] == k)
        lines.append(f"| {k} | {n} |")

    for folder, title, files in results:
        lines += ["", f"## {title} (`scenarios/{folder}.yaml`)", "",
                  "| 파일 | 행 | 열 이름 (기준명) | 오류 종류 | 원래 값 | 설명 |",
                  "|---|---|---|---|---|---|"]
        for f in files:
            for e in sorted(f.errors, key=lambda e: e["row"]):
                if e["kind"] == DUPLICATE:
                    col = f"(행 전체) — 중복기준: {e['key']}"
                else:
                    col = f"{e['col']} ({e['key']})"
                lines.append(
                    f"| {f.filename} | {e['row']} | {col} | {e['kind']} | {md_value(e['value'])} | {e['desc']} |"
                )

    (SAMPLES / "expected_errors.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return total


def json_value(v):
    if isinstance(v, date):
        return v.isoformat()
    return v


def write_answer_json(results):
    """테스트가 자동으로 대조하는 정답지. .md와 같은 기록에서 만든다."""
    data = {"version": 1, "note": FICTION_NOTE, "scenarios": {}}
    for folder, title, files in results:
        errors = []
        for f in files:
            for e in sorted(f.errors, key=lambda e: e["row"]):
                errors.append({
                    "file": f.filename,
                    "row": e["row"],
                    "column": e["col"],        # 파일에 실제로 쓰인 열 이름 (중복은 "(행 전체)")
                    "standard": e["key"],      # 기준명 (중복은 "품목코드 + 로트번호" 형태)
                    "kind": e["kind"],
                    "value": json_value(e["raw"]),
                    "shown": md_value(e["value"]),
                    "description": e["desc"],
                    "duplicate_of": e["duplicate_of"],
                })
        data["scenarios"][folder] = {
            "name": title,
            "folder": folder,
            "files": [{"file": f.filename, "header_row": f.header_row, "data_rows": len(f.rows)}
                      for f in files],
            "errors": errors,
        }
    (SAMPLES / "expected_errors.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main():
    rng = random.Random(20260928)  # 고정 시드: 몇 번을 돌려도 같은 샘플
    results = []
    for folder, title, build in SCENARIOS:
        out = SAMPLES / folder
        out.mkdir(parents=True, exist_ok=True)
        files = build(rng)
        for f in files:
            f.save(out)
            print(f"  {folder}/{f.filename}: 데이터 {len(f.rows)}행, 오류 {len(f.errors)}건")
        results.append((folder, title, files))
    total = write_answer_key(results)
    write_answer_json(results)
    print(f"정답지 작성: samples/expected_errors.md, samples/expected_errors.json (총 {total}건)")


if __name__ == "__main__":
    main()
