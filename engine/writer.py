"""결과 엑셀을 새 파일로 저장한다 (원본은 절대 건드리지 않는다).

시트 1 "취합결과": 출처 파일, 원래 행 + 기준열. 오류 셀은 오류 종류별 색.
시트 2 "오류목록": 파일, 행, 열, 기준열, 오류 종류, 값, 설명, 수정 제안값, 중복 그룹 + 색 범례.
"""

from __future__ import annotations

import re
import uuid
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .merge import MergeResult
from .validate import ALLOWED, DUPLICATE, FORMAT, KINDS, RANGE, REQUIRED, WHOLE_FILE

RESULT_SHEET = "취합결과"
ERROR_SHEET = "오류목록"
SOURCE_COLS = ["출처 파일", "원래 행"]
ERROR_COLS = ["파일", "행", "열", "기준열", "오류 종류", "값", "설명", "수정 제안값", "중복 그룹"]

COLORS = {
    REQUIRED: "FFC7CE",   # 연한 빨강
    FORMAT: "FFEB9C",     # 연한 노랑
    RANGE: "F8CBAD",      # 연한 주황
    ALLOWED: "D9D2E9",    # 연한 보라
    DUPLICATE: "BDD7EE",  # 연한 파랑
}
FILLS = {k: PatternFill("solid", start_color=v, end_color=v) for k, v in COLORS.items()}
BOLD = Font(bold=True)

# 오류목록 시트에 적는 최대 행 수. 넘으면 안내 행을 남긴다 (보안 검토 4번: 메모리 폭주 방지)
MAX_ERROR_ROWS = 20_000

_SAFE_NAME = re.compile(r"[^0-9A-Za-z가-힣_\-]+")


def _cell(ws, value: Any, fill: PatternFill | None = None, bold: bool = False) -> WriteOnlyCell:
    """write-only 시트용 셀. 문자열은 수식으로 해석되지 않게 항상 문자 타입으로 둔다."""
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        cell.data_type = "s"   # '=' '+' '-' '@'로 시작해도 수식이 되지 않는다
    elif isinstance(value, datetime):
        cell.number_format = "yyyy-mm-dd hh:mm:ss" if value.time() != datetime.min.time() else "yyyy-mm-dd"
    elif isinstance(value, date):
        cell.number_format = "yyyy-mm-dd"
    if fill is not None:
        cell.fill = fill
    if bold:
        cell.font = BOLD
    return cell


def safe_output_path(out_dir: str | Path, file_name: str, sources: list[Path]) -> Path:
    """output 폴더 안의 경로만 허용하고, 원본 파일이나 기존 파일과 겹치지 않는 새 파일을 만든다.

    파일은 배타적으로(없을 때만) 만들어 자리를 잡아 두므로, 동시에 실행돼도 서로 덮어쓰지 않는다.
    """
    out_dir = Path(out_dir).resolve()
    stem = _SAFE_NAME.sub("_", Path(file_name).stem).strip("_") or "취합결과"
    base = (out_dir / f"{stem}.xlsx").resolve()
    if base.parent != out_dir:
        raise ValueError("결과 파일은 출력 폴더 안에만 저장할 수 있습니다.")
    src = {Path(s).resolve() for s in sources}
    path, n = base, 2
    while True:
        if path not in src:
            try:
                with open(path, "xb"):
                    pass
                return path
            except FileExistsError:
                pass
        path = base.with_name(f"{base.stem}_{n}.xlsx")
        n += 1


def _set_widths(ws, widths: list[int]):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(8, min(w + 2, 60))


def _text_len(v: Any) -> int:
    if v is None:
        return 0
    s = v.isoformat() if isinstance(v, (date, datetime)) else str(v)
    return sum(2 if ord(ch) > 0x2E80 else 1 for ch in s)


def _result_value(row, name: str) -> Any:
    # 오류 셀은 원래 값 그대로 (자동으로 고치지 않는다), 정상 셀은 정규화된 값
    return row.raw[name] if name in row.errors else row.values[name]


def _error_values(issue) -> list[Any]:
    return [issue.file,
            issue.row if issue.row is not None else WHOLE_FILE,
            issue.column if issue.column is not None else "(없음)",
            issue.standard, issue.kind,
            issue.value if issue.value is not None else "(빈칸)",
            issue.message, issue.suggestion, issue.dup_group]


def write_result(result: MergeResult, out_dir: str | Path = "output",
                 file_name: str | None = None) -> Path:
    """결과 엑셀을 새 파일로 저장한다.

    write-only 모드로 한 행씩 흘려 쓰므로 행이 많아도 메모리를 적게 쓴다.
    열 너비는 행을 쓰기 전에 정해야 해서 먼저 한 번 훑어 계산한다.
    """
    scenario = result.scenario
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    if file_name is None:
        file_name = f"{scenario.key}_취합결과_{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:6]}"
    sources = [p.table.path for p in result.plans]
    path = safe_output_path(out_dir, file_name, sources)

    wb = Workbook(write_only=True)

    # ---- 시트 1: 취합결과
    ws = wb.create_sheet(RESULT_SHEET)
    headers = SOURCE_COLS + scenario.column_names
    widths = [_text_len(h) for h in headers]
    for row in result.rows:
        widths[0] = max(widths[0], _text_len(row.source_file))
        for c, name in enumerate(scenario.column_names, start=2):
            widths[c] = max(widths[c], _text_len(_result_value(row, name)))
    _set_widths(ws, widths)
    ws.freeze_panes = "C2"
    if result.rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(result.rows) + 1}"
    ws.append([_cell(ws, h, bold=True) for h in headers])
    for row in result.rows:
        dup_fill = FILLS[DUPLICATE] if row.dup_group else None
        cells = [_cell(ws, row.source_file, dup_fill), _cell(ws, row.excel_row, dup_fill)]
        for name in scenario.column_names:
            kind = row.errors.get(name)
            if kind:
                fill = FILLS[kind]
            elif row.dup_group and name in scenario.dup_keys:
                fill = FILLS[DUPLICATE]
            else:
                fill = None
            cells.append(_cell(ws, _result_value(row, name), fill))
        ws.append(cells)

    # ---- 시트 2: 오류목록 (+ 오른쪽에 색 범례)
    es = wb.create_sheet(ERROR_SHEET)
    shown = result.issues[:MAX_ERROR_ROWS]
    omitted = len(result.issues) - len(shown)
    widths = [_text_len(h) for h in ERROR_COLS]
    for issue in shown:
        for c, v in enumerate(_error_values(issue)):
            widths[c] = max(widths[c], _text_len(v))
    _set_widths(es, widths + [0, 16])
    es.freeze_panes = "A2"
    if shown:
        es.auto_filter.ref = f"A1:{get_column_letter(len(ERROR_COLS))}{len(shown) + 1}"

    legend = [_cell(es, "색 범례", bold=True)] + [_cell(es, k, FILLS[k]) for k in KINDS] +              [None, _cell(es, "취합결과 시트의 오류 셀은 원래 값 그대로 두고 색만 칠했습니다.")]
    kind_col = ERROR_COLS.index("오류 종류")

    def with_legend(cells: list, i: int) -> list:
        extra = legend[i] if i < len(legend) else None
        return cells + [None, extra] if extra is not None else cells

    es.append(with_legend([_cell(es, h, bold=True) for h in ERROR_COLS], 0))
    line = 1
    for issue in shown:
        cells = [_cell(es, v, FILLS[issue.kind] if c == kind_col else None)
                 for c, v in enumerate(_error_values(issue))]
        es.append(with_legend(cells, line))
        line += 1
    if omitted:
        es.append(with_legend([_cell(es, f"이하 {omitted:,}건은 생략했습니다 (최대 {MAX_ERROR_ROWS:,}건 표시). "
                                          "원본 파일을 먼저 정리한 뒤 다시 취합해 주세요.")], line))
        line += 1
    while line < len(legend):   # 오류가 적어 범례가 더 길면 나머지 범례 행
        es.append(with_legend([None] * len(ERROR_COLS), line))
        line += 1

    wb.save(path)
    return path
