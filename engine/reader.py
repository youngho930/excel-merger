"""엑셀 파일을 읽고 머리글 행을 자동으로 찾는다.

원본 파일은 읽기 전용으로만 연다 (절대 저장하지 않는다).
행 번호는 사용자가 엑셀 화면에서 보는 번호(1부터)를 그대로 쓴다.
"""

from __future__ import annotations

import posixpath
import zipfile
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .normalize import is_blank, normalize_header
from .scenario import Scenario

HEADER_SCAN_ROWS = 20           # 머리글을 찾을 때 위에서부터 훑는 행 수
EXCEL_SUFFIXES = (".xlsx", ".xlsm")

# 자원 한도: 업로드된 파일 하나가 메모리·CPU를 독차지하지 못하게 한다 (보안 검토 1번)
MAX_FILE_BYTES = 20 * 1024 * 1024          # 압축된 파일 크기
MAX_UNZIPPED_BYTES = 100 * 1024 * 1024     # 압축을 푼 전체 크기
MAX_SHARED_STRINGS_BYTES = 30 * 1024 * 1024
MAX_ZIP_ENTRIES = 1000
MAX_COMPRESSION_RATIO = 500                # 1MB 넘는 항목의 압축률 상한 (0으로 채운 압축 폭탄은 약 1000배, 반복이 많은 정상 데이터는 약 230배)
MAX_SHEETS = 20
MAX_ROWS = 50_000                          # 파일 하나에서 읽는 데이터 최대 행 수
MAX_COLS = 200
MAX_CELLS = 1_000_000                      # 파일 하나에서 읽는 최대 칸 수 (행 × 열)
MAX_SHEET_XML_BYTES = 60 * 1024 * 1024     # 모든 시트가 가리키는 시트 XML 크기의 합

_NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


class ReadError(Exception):
    """엑셀을 읽을 수 없을 때. 메시지는 한국어."""


@dataclass
class SourceRow:
    excel_row: int              # 엑셀 화면의 행 번호
    cells: tuple[Any, ...]      # 열 위치(0부터) 순서의 원래 값


@dataclass
class SourceTable:
    file_name: str
    path: Path
    sheet_name: str
    header_row: int                     # 머리글 행 번호 (엑셀 화면 기준)
    headers: list[str]                  # 표시용 열 이름 (빈 머리글은 "(빈 머리글 C열)")
    raw_headers: list[Any]              # 원래 머리글 값 (매칭에 사용)
    rows: list[SourceRow] = field(default_factory=list)
    header_confidence: str = "높음"     # "높음" | "낮음" | "지정"

    def column_label(self, index: int) -> str:
        """'C열 'Part No'' 같은 표시."""
        return f"{get_column_letter(index + 1)}열 '{self.headers[index]}'"


@dataclass(frozen=True)
class HeaderGuess:
    row: int            # 1부터
    score: int          # 기준명·동의어와 일치한 서로 다른 기준열 수
    confident: bool


def header_vocabulary(scenario: Scenario) -> dict[str, str]:
    """정규화한 이름 -> 기준명."""
    vocab: dict[str, str] = {}
    for col in scenario.columns:
        for n in col.names:
            vocab[normalize_header(n)] = col.name
    return vocab


def detect_header_row(rows: list[tuple[Any, ...]], scenario: Scenario) -> HeaderGuess:
    """위쪽 행들 중 기준명·동의어와 가장 많이 일치하는 행을 머리글로 본다."""
    vocab = header_vocabulary(scenario)
    threshold = min(2, len(scenario.columns))
    best_row, best_score = 0, -1
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        matched = {vocab[k] for k in (normalize_header(v) for v in row if not is_blank(v)) if k in vocab}
        if len(matched) > best_score:
            best_row, best_score = i, len(matched)
    if best_score >= threshold:
        return HeaderGuess(best_row + 1, best_score, True)
    # 찾지 못함: 채워진 칸이 가장 많은 행 (동점이면 위쪽)
    fallback, most = 0, -1
    for i, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        n = sum(1 for v in row if not is_blank(v))
        if n > most:
            fallback, most = i, n
    return HeaderGuess(fallback + 1, max(best_score, 0), False)


def _check_zip(path: Path) -> None:
    """openpyxl로 열기 전에 압축을 푼 크기를 검사한다 (작은 파일이 풀리면서 커지는 압축 폭탄 방지)."""
    too_big = f"'{path.name}'은 압축을 풀면 너무 커서 처리할 수 없습니다."
    try:
        with zipfile.ZipFile(path) as zf:
            infos = zf.infolist()
    except (zipfile.BadZipFile, OSError, ValueError):
        raise ReadError(f"'{path.name}'을 엑셀 파일로 열 수 없습니다. 파일이 손상되었거나 암호가 걸려 있을 수 있습니다.") from None
    if len(infos) > MAX_ZIP_ENTRIES:
        raise ReadError(too_big)
    total = 0
    for info in infos:
        total += info.file_size
        if total > MAX_UNZIPPED_BYTES:
            raise ReadError(too_big)
        if info.file_size > 1024 * 1024 and info.file_size > MAX_COMPRESSION_RATIO * max(info.compress_size, 1):
            raise ReadError(too_big)
        if info.filename.lower().endswith("sharedstrings.xml") and info.file_size > MAX_SHARED_STRINGS_BYTES:
            raise ReadError(too_big)

    # openpyxl은 파일을 열 때 시트마다 시트 XML을 훑는다 (크기 정보가 없으면 끝까지).
    # 여러 시트가 같은 XML을 가리키면 같은 내용을 여러 번 읽으므로, 시트가 가리키는 XML 크기의
    # 합에 상한을 둔다 (60MB 시트 하나를 시트 4개가 공유하면 여는 데만 50초 걸렸음).
    sizes = {i.filename: i.file_size for i in infos}
    try:
        with zipfile.ZipFile(path) as zf:
            targets = _sheet_targets(zf, sizes)
    except ReadError:
        raise
    except Exception:
        targets = None   # 구조를 알 수 없으면 openpyxl이 열 때 판단하게 둔다
    if targets is not None:
        if len(targets) > MAX_SHEETS:
            raise ReadError(f"'{path.name}'에 시트가 너무 많습니다 (최대 {MAX_SHEETS}개).")
        if sum(sizes.get(t, 0) for t in targets) > MAX_SHEET_XML_BYTES:
            raise ReadError(too_big)


_REL_NS = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_DOC_REL_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
MAX_META_XML_BYTES = 1024 * 1024   # workbook.xml·관계 파일은 작다. 크면 거부


def _read_small_xml(zf: zipfile.ZipFile, name: str, sizes: dict[str, int]):
    if sizes.get(name, 0) > MAX_META_XML_BYTES:
        raise ReadError("엑셀 파일 구조 정보가 비정상적으로 큽니다.")
    return ET.fromstring(zf.read(name))


def _resolve(base_dir: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    return posixpath.normpath(posixpath.join(base_dir, target))


def _sheet_targets(zf: zipfile.ZipFile, sizes: dict[str, int]) -> list[str]:
    """workbook.xml의 시트 목록을 시트 XML 경로 목록으로 바꾼다 (시트마다 1개, 중복 포함)."""
    root_rels = _read_small_xml(zf, "_rels/.rels", sizes)
    wb_path = next(_resolve("", r.get("Target", "")) for r in root_rels.iter(_REL_NS + "Relationship")
                   if r.get("Type", "").endswith("/officeDocument"))
    wb_dir = posixpath.dirname(wb_path)
    rels = _read_small_xml(zf, posixpath.join(wb_dir, "_rels", posixpath.basename(wb_path) + ".rels"), sizes)
    by_id = {r.get("Id"): _resolve(wb_dir, r.get("Target", "")) for r in rels.iter(_REL_NS + "Relationship")}
    wb = _read_small_xml(zf, wb_path, sizes)
    return [by_id.get(s.get(_DOC_REL_NS + "id"), "") for s in wb.iter(_NS_MAIN + "sheet")]


def _open_workbook(path: Path):
    if path.suffix.lower() not in EXCEL_SUFFIXES:
        raise ReadError(f"'{path.name}'은 엑셀(.xlsx) 파일이 아닙니다. .xls 파일은 엑셀에서 .xlsx로 다시 저장해 주세요.")
    if not path.is_file():
        raise ReadError(f"'{path.name}' 파일을 찾을 수 없습니다.")
    if path.stat().st_size > MAX_FILE_BYTES:
        raise ReadError(f"'{path.name}'이 너무 큽니다 (최대 {MAX_FILE_BYTES // (1024 * 1024)}MB).")
    _check_zip(path)
    try:
        # read_only: 원본을 수정할 수 없고 메모리도 적게 쓴다. data_only: 수식 대신 저장된 계산값.
        return load_workbook(path, read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError, TypeError):
        raise ReadError(f"'{path.name}'을 엑셀 파일로 열 수 없습니다. 파일이 손상되었거나 암호가 걸려 있을 수 있습니다.") from None
    except MemoryError:
        raise
    except Exception:  # openpyxl 내부 예외 종류가 다양하다
        raise ReadError(f"'{path.name}'을 엑셀 파일로 열 수 없습니다.") from None


def _blank_row(row) -> bool:
    return all(is_blank(v) for v in row)


def _trim(row) -> tuple[Any, ...]:
    """뒤쪽 빈칸을 잘라낸 행."""
    row = tuple(row)
    end = len(row)
    while end and is_blank(row[end - 1]):
        end -= 1
    return row[:end]


def _scan_top(ws) -> list[tuple[Any, ...]]:
    """머리글 탐지용으로 위쪽 HEADER_SCAN_ROWS 행만 읽는다.

    max_row·max_col을 주면 openpyxl이 그 밖의 행·열을 만들지 않는다
    (파일에 적힌 행 번호가 10억이어도 바로 멈춘다).
    """
    ws.reset_dimensions()  # 파일에 적힌 크기 정보를 믿지 않는다
    rows = [_trim(r) for r in ws.iter_rows(max_row=HEADER_SCAN_ROWS, max_col=MAX_COLS + 1, values_only=True)]
    if any(len(r) > MAX_COLS for r in rows):
        raise ReadError(f"'{ws.title}' 시트의 열이 너무 많습니다 (최대 {MAX_COLS}열).")
    return rows


def _has_data_after(path: Path, ws, last_row: int) -> bool:
    """last_row보다 아래에 값이 있는 칸이 하나라도 있는지 시트 XML을 훑어 확인한다.

    iter_rows(max_row=…)는 한도 밖의 행을 조용히 건너뛰므로, 한도를 넘는 데이터가
    말없이 빠지지 않도록 따로 확인한다. 빈 행을 만들어 내지 않으므로 행 번호가
    아무리 멀리 떨어져 있어도 빠르다. (시트 XML 경로는 openpyxl read-only 시트의 내부 속성)
    """
    member = getattr(ws, "_worksheet_path", None)
    if not member:
        return False
    row_no = 0
    with zipfile.ZipFile(path) as zf, zf.open(member.lstrip("/")) as src:
        for _, el in ET.iterparse(src, events=("end",)):
            if el.tag != _NS_MAIN + "row":
                continue
            r = el.get("r")
            row_no = int(r) if r and r.isdigit() else row_no + 1
            if row_no > last_row and any(
                    c.find(_NS_MAIN + "v") is not None or c.find(_NS_MAIN + "is") is not None
                    for c in el.iter(_NS_MAIN + "c")):
                return True
            el.clear()
    return False


def _read_data(ws, header_row: int, width: int, path: Path, max_rows: int = MAX_ROWS) -> list[SourceRow]:
    """머리글 다음 행부터 읽는다. 열은 width까지만, 행·칸 수는 한도까지만."""
    width = max(width, 1)
    last_allowed = header_row + max_rows
    if max_rows * width > MAX_CELLS:
        last_allowed = header_row + MAX_CELLS // width
    out: list[SourceRow] = []
    for offset, row in enumerate(ws.iter_rows(min_row=header_row + 1, max_row=last_allowed,
                                              max_col=width, values_only=True)):
        if _blank_row(row):
            continue  # 완전히 빈 행은 건너뛴다
        out.append(SourceRow(excel_row=header_row + 1 + offset,
                             cells=tuple(row) + (None,) * (width - len(row))))
    if _has_data_after(path, ws, last_allowed):
        limit = min(max_rows, MAX_CELLS // width)
        raise ReadError(f"'{ws.title}' 시트의 데이터가 너무 많습니다 (열 {width}개 기준 최대 {limit:,}행).")
    return out


def read_table(path: str | Path, scenario: Scenario, header_row: int | None = None,
               sheet: str | None = None, max_rows: int | None = None) -> SourceTable:
    """엑셀 파일 하나를 읽는다.

    max_rows: 데이터 최대 행 수 (배포 환경별 제한). MAX_ROWS보다 크게 할 수는 없다.
    header_row를 주면 그 행을 머리글로 쓰고, 없으면 자동으로 찾는다.
    sheet를 주지 않으면 머리글 점수가 가장 높은 시트를 쓴다(동점이면 앞 시트).
    머리글 탐지는 시트마다 위쪽 몇 행만 읽고, 고른 시트 하나만 끝까지 읽는다.
    """
    path = Path(path)
    wb = _open_workbook(path)
    try:
        if len(wb.sheetnames) > MAX_SHEETS:
            raise ReadError(f"'{path.name}'에 시트가 너무 많습니다 (최대 {MAX_SHEETS}개).")
        names = [sheet] if sheet else wb.sheetnames
        best = None
        for name in names:
            if name not in wb.sheetnames:
                raise ReadError(f"'{path.name}'에 '{name}' 시트가 없습니다.")
            top = _scan_top(wb[name])
            guess = detect_header_row(top, scenario)
            if best is None or (guess.confident, guess.score) > (best[0].confident, best[0].score):
                best = (guess, name, top)
        guess, sheet_name, top = best
        ws = wb[sheet_name]

        if header_row is not None:
            if header_row < 1 or header_row > HEADER_SCAN_ROWS:
                raise ReadError(f"머리글 행은 1~{HEADER_SCAN_ROWS}행 중에서 지정해 주세요.")
            hr, confidence = header_row, "지정"
        else:
            hr, confidence = guess.row, ("높음" if guess.confident else "낮음")

        # 열 너비: 위쪽에서 본 가장 넓은 행 기준 (그보다 오른쪽 칸은 머리글이 없어 읽지 않는다)
        width = max((len(r) for r in top), default=0)
        rows = _read_data(ws, hr, width, path, min(max_rows or MAX_ROWS, MAX_ROWS))
    except (ReadError, MemoryError):   # 메모리 부족은 "손상된 파일"이 아니다 (jobs 가 따로 안내)
        raise
    except Exception:
        # read_only 모드는 시트를 읽을 때 비로소 XML을 해석한다. 손상된 파일의 내부 예외를
        # 그대로 내보내지 않는다 (내부 경로·스택 노출 방지, 보안 검토 5번)
        raise ReadError(f"'{path.name}'의 시트 내용을 읽을 수 없습니다. 파일이 손상되었을 수 있습니다.") from None
    finally:
        wb.close()

    header_vals = list(top[hr - 1]) if hr - 1 < len(top) else []
    raw_headers = header_vals + [None] * (width - len(header_vals))
    headers = [
        str(v).strip() if not is_blank(v) else f"(빈 머리글 {get_column_letter(i + 1)}열)"
        for i, v in enumerate(raw_headers)
    ]
    return SourceTable(file_name=path.name, path=path, sheet_name=sheet_name, header_row=hr,
                       headers=headers, raw_headers=raw_headers, rows=rows,
                       header_confidence=confidence)
