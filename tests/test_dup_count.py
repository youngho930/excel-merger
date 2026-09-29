"""중복 건수 표기: 오류 건수는 행 기준, 정답지는 쌍 기준이라 '6행 (3쌍)'처럼 함께 보여준다."""

from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from answer_key import ROOT, compare, load_free_answer
from engine.free_form import regroup
from engine.merge import MergeResult
from engine.validate import DUPLICATE, REQUIRED, Issue

APP = str(ROOT / "app.py")
ANSWER = load_free_answer()
PAIRS = [("고객명", "성명"), ("연락처", "전화번호"), ("제품명", "모델"), ("증상", "고장내용"), ("접수일", "등록일")]


def issue(kind, group=None, row=2):
    return Issue("a.xlsx", row, "(행 전체)", "키", kind, None, "", dup_group=group)


def result_of(*issues):
    return MergeResult(scenario=None, plans=[], rows=[], issues=list(issues))


def test_dup_count_text_pairs_and_groups():
    pairs = result_of(issue(DUPLICATE, "중복-1"), issue(DUPLICATE, "중복-1", 3),
                      issue(DUPLICATE, "중복-2", 4), issue(DUPLICATE, "중복-2", 5), issue(REQUIRED))
    assert pairs.dup_count_text() == "4행 (2쌍)"                   # 필수값 빈칸은 세지 않는다
    triple = result_of(*[issue(DUPLICATE, "중복-1", r) for r in (2, 3, 4)], issue(DUPLICATE, "중복-2", 5),
                       issue(DUPLICATE, "중복-2", 6))
    assert triple.dup_count_text() == "5행 (2그룹)"                # 세 행짜리 그룹이 있으면 '쌍'이라 하지 않는다
    assert result_of(issue(REQUIRED)).dup_count_text() == "0행 (0쌍)"


def run_free_demo() -> AppTest:
    """자유 양식 체험: 미리 설정된 필수·중복 기준 그대로 5쌍만 묶고 실행 (정답지 8건)."""
    at = AppTest.from_file(APP, default_timeout=120)
    at.run()
    at.button(key="free_demo_btn").click().run()
    for a, b in PAIRS:
        g = {x.label: x.gid for x in regroup(at.session_state["ff_tables"], at.session_state["ff_merges"])}
        at.multiselect(key=f"ff_merge_pick_{at.session_state['load_id']}").set_value([g[a], g[b]])
        at.button(key="ff_merge_btn").click().run()
    at.button(key="run_btn").click().run()
    assert not at.exception, at.exception
    assert compare(at.session_state["result"].issues, ANSWER["errors"]).ok
    return at


def test_screen_and_excel_show_rows_and_pairs():
    at = run_free_demo()
    answer_dups = sum(1 for e in ANSWER["errors"] if e["kind"] == DUPLICATE)
    assert answer_dups == 3                                          # 정답지는 쌍 기준 3건
    metrics = {m.label: m.value for m in at.metric}
    assert metrics["중복 행"] == "6행 (3쌍)"
    total = len(at.session_state["result"].issues)
    assert metrics["전체 오류"] == f"{total}건" and total == len(ANSWER["errors"]) + answer_dups   # 행 기준이라 3건 더 많다
    assert "행 기준" in [c.value for c in at.caption]                 # 전체 오류 옆 작은 표시

    wb = load_workbook(Path(at.session_state["current_output"]))
    rows = [[c.value for c in r] for r in wb["요약"].iter_rows()]
    head = next(i for i, r in enumerate(rows) if r[0] == "오류 종류")
    assert rows[head][4] == "비고"
    top = {r[0]: r[1] for r in rows[:head] if r and r[0]}
    assert top["중복 행"] == "6행 (3쌍)"
    body = {r[0]: r for r in rows[head + 1:] if r and r[0]}
    assert body[DUPLICATE][1] == 6 and body[DUPLICATE][4] == "6행 (3쌍)"
    assert body["합계"][4].startswith("행 기준")
