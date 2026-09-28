"""보안 검토(AI-Generated Code Security Auditor)에서 지적된 문제의 회귀 테스트."""

import time
import zipfile
from datetime import date

import pytest
from openpyxl import Workbook, load_workbook

from engine import ReadError, execute, load_scenario, prepare, read_table, write_result
from engine import writer as writer_mod
from engine.merge import MergedRow, _find_duplicates, escape_formula
from engine.normalize import parse_date
from engine.scenario import ScenarioError, find_scenario

from answer_key import ROOT

YAML = """name: 보안 테스트
columns:
  - 기준명: 코드
    필수: true
    형식: 문자
  - 기준명: 날짜
    형식: 날짜
중복기준: [코드]
"""


@pytest.fixture
def sc(tmp_path):
    p = tmp_path / "sec.yaml"
    p.write_text(YAML, encoding="utf-8")
    return load_scenario(p)


def xlsx(path, rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def replace_in_zip(path, member, new_bytes):
    """xlsx(zip) 안의 파일 하나를 바꿔 손상된 파일을 만든다."""
    with zipfile.ZipFile(path) as zf:
        items = {i.filename: zf.read(i.filename) for i in zf.infolist()}
    items[member] = new_bytes
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items.items():
            zf.writestr(name, data)


def test_date_regex_is_not_quadratic():
    # 2번: 공백 32,000개가 들어간 날짜 문자열 (수정 전 9.45초)
    t = time.perf_counter()
    p = parse_date("2024-01-01" + " " * 32000 + "x")
    assert not p.ok and time.perf_counter() - t < 0.1
    assert parse_date("2026-09-03 12:30").value == date(2026, 9, 3)
    assert parse_date("2026-09-03T12:30:00").value == date(2026, 9, 3)


def test_duplicate_descriptions_are_linear(sc):
    # 3번: 같은 키를 가진 행 20,000개 (수정 전 9.6초)
    rows = [MergedRow("a.xlsx", i + 2, {"코드": "X", "날짜": None}, {}) for i in range(20000)]
    t = time.perf_counter()
    issues = _find_duplicates(sc, [(r, ("X",)) for r in rows], rows)
    assert time.perf_counter() - t < 2
    assert len(issues) == 20000
    assert issues[0].message.startswith("a.xlsx 3행, a.xlsx 4행") and "외 19994행" in issues[0].message


def test_far_row_number_is_rejected_quickly(sc, tmp_path):
    # 1번: 멀리 떨어진 행 번호 하나 (수정 전 행 수만큼 빈 행을 만들며 멈춤)
    wb = Workbook()
    ws = wb.active
    ws.append(["코드", "날짜"])
    ws.cell(row=1_000_000, column=1, value="x")
    wb.save(tmp_path / "far.xlsx")
    t = time.perf_counter()
    with pytest.raises(ReadError, match="너무 많습니다"):
        read_table(tmp_path / "far.xlsx", sc)
    assert time.perf_counter() - t < 5


def test_wide_empty_cells_do_not_widen_rows(sc, tmp_path):
    # 1번: 머리글 오른쪽 먼 열(XFD)의 빈 서식 셀
    wb = Workbook()
    ws = wb.active
    ws.append(["코드", "날짜"])
    for r in range(2, 2002):
        ws.cell(row=r, column=1, value=f"C{r}")
        ws.cell(row=r, column=16384).number_format = "0"
    wb.save(tmp_path / "wide.xlsx")
    t = read_table(tmp_path / "wide.xlsx", sc)
    assert len(t.headers) == 2 and len(t.rows) == 2000


def test_sheet_structure_is_read_for_normal_files(tmp_path):
    # 1번: 시트 XML 크기 합 검사가 정상 파일의 시트 구조를 제대로 읽는지 (못 읽으면 검사가 조용히 빠진다)
    from engine.reader import _sheet_targets
    wb = Workbook()
    wb.create_sheet("둘째")
    wb.create_sheet("셋째")
    wb.save(tmp_path / "multi.xlsx")
    with zipfile.ZipFile(tmp_path / "multi.xlsx") as zf:
        sizes = {i.filename: i.file_size for i in zf.infolist()}
        targets = _sheet_targets(zf, sizes)
    assert len(targets) == 3 and all(t in sizes for t in targets)
    for sample in (ROOT / "samples").glob("*/*.xlsx"):
        with zipfile.ZipFile(sample) as zf:
            sizes = {i.filename: i.file_size for i in zf.infolist()}
            assert all(t in sizes for t in _sheet_targets(zf, sizes))


def test_zip_bomb_rejected(sc, tmp_path):
    # 1번: 압축률이 비정상적으로 높은 항목 (2MB의 0 → 수 KB)
    p = xlsx(tmp_path / "bomb.xlsx", [["코드", "날짜"], ["A", None]])
    replace_in_zip(p, "xl/media/zeros.bin", b"\0" * (2 * 1024 * 1024))
    with pytest.raises(ReadError, match="압축을 풀면 너무 커서"):
        read_table(p, sc)


@pytest.mark.parametrize("cell_xml", [
    '<c r="A2"><v>abc</v></c>',          # 숫자 칸에 글자
    '<c r="A2" t="s"><v>999</v></c>',    # 없는 공유 문자열 번호
    '<c r="A2" t="d"><v>zzz</v></c>',    # 잘못된 날짜
])
def test_corrupt_sheet_gives_korean_read_error(sc, tmp_path, cell_xml):
    # 5번: 시트 XML 손상 시 내부 예외 대신 한국어 ReadError
    p = xlsx(tmp_path / "bad.xlsx", [["코드", "날짜"]])
    sheet = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
             '<row r="1"><c r="A1" t="inlineStr"><is><t>코드</t></is></c>'
             '<c r="B1" t="inlineStr"><is><t>날짜</t></is></c></row>'
             f'<row r="2">{cell_xml}</row></sheetData></worksheet>').encode("utf-8")
    replace_in_zip(p, "xl/worksheets/sheet1.xml", sheet)
    with pytest.raises(ReadError, match="손상"):
        read_table(p, sc)


def test_cli_hides_internal_errors(tmp_path, capsys, monkeypatch):
    # 5번: CLI가 예상하지 못한 예외에서도 트레이스백·내부 경로를 보여주지 않는다
    from answer_key import run_merge_main
    import engine.merge as merge_mod

    def boom(*a, **k):
        raise RuntimeError(r"C:\secret\path")
    monkeypatch.setattr(merge_mod, "read_table", boom)
    code, out = run_merge_main(["stock_count", str(ROOT / "samples" / "stock_count"),
                                "--out", str(tmp_path)])
    err = capsys.readouterr().err
    assert code == 1 and "secret" not in out + err and "Traceback" not in err


def test_yaml_alias_bomb_rejected(tmp_path):
    # 6번: 동의어 안에 앵커(&)·별칭(*)으로 중첩 목록을 만든다 (수정 전 479바이트 → 913MB)
    parts = ["&a0 [" + ", ".join(["xyz"] * 10) + "]"]
    for i in range(1, 7):
        parts.append(f"&a{i} [" + ", ".join([f"*a{i - 1}"] * 10) + "]")
    text = "\n".join([
        "name: 폭탄",
        "columns:",
        "  - 기준명: 코드",
        "    형식: 문자",
        "    동의어: [" + ", ".join(parts) + "]",
        "중복기준: []",
        "",
    ])
    p = tmp_path / "bomb.yaml"
    p.write_text(text, encoding="utf-8")
    assert p.stat().st_size < 1000
    t = time.perf_counter()
    with pytest.raises(ScenarioError, match="목록 안에 목록 불가"):
        load_scenario(p)
    assert time.perf_counter() - t < 1


def test_scenario_name_with_trailing_newline_rejected():
    # 7번
    with pytest.raises(ScenarioError, match="영문"):
        find_scenario("stock_count\n", ROOT / "scenarios")


def test_formula_values_stay_text_in_result(sc, tmp_path):
    # 원본에 수식처럼 보이는 "글자"가 들어 있는 경우 (openpyxl은 '='로 시작하면 수식으로 저장하므로 문자로 고정)
    p = xlsx(tmp_path / "f.xlsx", [["코드", "날짜"], ["=1+1", "=HYPERLINK(1)"], ["@SUM(1,1)", None]])
    wb = load_workbook(p)
    for row in wb.active.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, str):
                c.data_type = "s"
    wb.save(p)
    result = execute(sc, prepare(sc, [p]))
    out = write_result(result, tmp_path / "out")
    with zipfile.ZipFile(out) as zf:
        sheets = [zf.read(n) for n in zf.namelist() if n.startswith("xl/worksheets/")]
    assert all(b"<f>" not in s for s in sheets)
    ws = load_workbook(out)["취합결과"]
    assert ws.cell(2, 3).value == "=1+1" and ws.cell(2, 3).data_type == "s"


def test_formula_escape_for_csv():
    # 8번
    assert escape_formula("=cmd") == "'=cmd" and escape_formula("-1") == "'-1"
    assert escape_formula("정상") == "정상" and escape_formula(5) == 5


def test_error_sheet_is_capped(sc, tmp_path, monkeypatch):
    # 4번: 오류목록 행 수 상한
    monkeypatch.setattr(writer_mod, "MAX_ERROR_ROWS", 3)
    p = xlsx(tmp_path / "many.xlsx", [["코드", "날짜"]] + [[None, "x"] for _ in range(5)])
    result = execute(sc, prepare(sc, [p]))
    assert len(result.issues) == 10
    es = load_workbook(write_result(result, tmp_path / "out"))["오류목록"]
    col_a = [es.cell(r, 1).value for r in range(2, es.max_row + 1) if es.cell(r, 1).value]
    assert len(col_a) == 4 and col_a[-1].startswith("이하 7건은 생략")


def test_output_names_never_collide(sc, tmp_path):
    # 9번: 같은 이름으로 연달아 저장해도 덮어쓰지 않는다
    p = xlsx(tmp_path / "a.xlsx", [["코드", "날짜"], ["A", None]])
    result = execute(sc, prepare(sc, [p]))
    a = write_result(result, tmp_path / "out", file_name="같은이름")
    b = write_result(result, tmp_path / "out", file_name="같은이름")
    assert a != b and a.exists() and b.exists()
