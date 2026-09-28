"""샘플에 없는 상황을 임시 YAML·임시 엑셀로 만들어 엔진 전체 흐름을 확인한다.

- 새 시나리오: 엔진 코드 수정 없이 YAML만 추가해도 동작하는지
- 필수 열이 아예 없는 파일
- 같은 기준열에 원본 열 후보가 둘 이상인 파일
- 날짜 칸의 숫자 셀 수정 제안이 결과 엑셀까지 나가는지
"""

from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from engine import execute, load_scenario, pre_run_warnings, prepare, write_result
from answer_key import run_merge_main

# 기존 시나리오와 겹치는 이름이 하나도 없는 가상 시나리오
NEW_YAML = """name: 설비 일일점검
description: 테스트용 가상 시나리오
columns:
  - 기준명: 설비ID
    동의어: [장비번호, Equip ID]
    필수: true
    형식: 문자
  - 기준명: 점검일
    동의어: [점검일자, Check Date]
    필수: true
    형식: 날짜
  - 기준명: 온도
    동의어: [온도(℃), Temp]
    필수: true
    형식: 실수
    범위: [-20, 80]
  - 기준명: 상태
    동의어: [설비상태, Status]
    필수: true
    형식: 문자
    허용값: [정상, 점검필요, 고장]
  - 기준명: 점검자
    동의어: [담당, Inspector]
    필수: false
    형식: 문자
중복기준: [설비ID, 점검일]
"""


def write_yaml(path, text=NEW_YAML):
    path.write_text(text, encoding="utf-8")
    return path


def make_xlsx(path, rows, title_lines=()):
    wb = Workbook()
    ws = wb.active
    for line in title_lines:
        ws.append([line])
    for r in rows:
        ws.append(list(r))
    wb.save(path)
    return path


def issue_keys(result):
    return {(i.file, i.row, i.standard, i.kind) for i in result.issues}


def test_new_scenario_works_without_code_change(tmp_path):
    sc = load_scenario(write_yaml(tmp_path / "equipment_check.yaml"))
    data = tmp_path / "data"
    data.mkdir()
    make_xlsx(data / "A공장.xlsx", [
        ("설비ID", "점검일", "온도", "상태", "점검자"),
        ("EQ-01", datetime(2026, 9, 1), 35.5, "정상", "김"),      # 2행
        ("EQ-02", datetime(2026, 9, 1), 120, "정상", "김"),       # 3행: 범위 밖
        (None, datetime(2026, 9, 1), 30, "정상", "김"),           # 4행: 필수값 빈칸
        ("EQ-04", datetime(2026, 9, 1), 31, "양호", None),        # 5행: 허용값 아님 (점검자는 선택)
    ])
    make_xlsx(data / "B공장.xlsx", [
        ("Status", "Temp", "Check Date", "Equip ID"),             # 다른 이름·순서, 점검자 열 없음
        ("정상", "28도", "2026.9.1", "EQ-10"),                    # 4행: 형식 오류
        ("점검필요", 40, "2026년 9월 1일", "EQ-01"),              # 5행: A공장 2행과 중복
    ], title_lines=("B공장 설비 일일점검표", "작성: 설비팀"))

    plans = prepare(sc, sorted(data.glob("*.xlsx")))
    b = next(p for p in plans if p.file_name == "B공장.xlsx")
    assert b.table.header_row == 3
    assert {m.standard: m.method for m in b.match.matches} == {
        "설비ID": "동의어", "점검일": "동의어", "온도": "동의어", "상태": "동의어", "점검자": "매칭 안 됨"}
    assert pre_run_warnings(plans) == []   # 빠진 열은 선택 열뿐

    result = execute(sc, plans)
    assert issue_keys(result) == {
        ("A공장.xlsx", 3, "온도", "범위 밖 값"),
        ("A공장.xlsx", 4, "설비ID", "필수값 빈칸"),
        ("A공장.xlsx", 5, "상태", "허용값 아닌 값"),
        ("B공장.xlsx", 4, "온도", "형식 오류"),
        ("A공장.xlsx", 2, "설비ID + 점검일", "중복 행"),
        ("B공장.xlsx", 5, "설비ID + 점검일", "중복 행"),
    }

    # CLI도 코드 수정 없이 YAML 경로만으로 동작
    code, out = run_merge_main([str(tmp_path / "equipment_check.yaml"), str(data), "--out", str(tmp_path / "out")])
    assert code == 0
    assert "오류 6건: 필수값 빈칸 1, 형식 오류 1, 범위 밖 값 1, 허용값 아닌 값 1, 중복 행 2" in out


def _stock_like(tmp_path):
    sc = load_scenario(write_yaml(tmp_path / "s.yaml"))
    ok = make_xlsx(tmp_path / "정상.xlsx", [
        ("설비ID", "점검일", "온도", "상태"),
        ("EQ-01", datetime(2026, 9, 1), 30, "정상"),
    ])
    return sc, ok


def test_missing_required_column(tmp_path):
    sc, ok = _stock_like(tmp_path)
    # 필수이면서 중복기준인 '설비ID' 열이 없다. '기계코드'는 어느 동의어에도 없는 이름.
    bad = make_xlsx(tmp_path / "열빠짐.xlsx", [
        ("기계코드", "점검일", "온도", "상태"),
        ("EQ-01", datetime(2026, 9, 1), 30, "정상"),   # ok 파일과 같은 값이지만 중복 검사 제외
        ("EQ-02", datetime(2026, 9, 2), 31, "정상"),
    ])
    plans = prepare(sc, [ok, bad])
    warnings = pre_run_warnings(plans)
    assert len(warnings) == 1
    assert "[열빠짐.xlsx] 필수 열 '설비ID'" in warnings[0]
    assert "매칭 안 된 원본 열: A열 '기계코드'" in warnings[0]

    result = execute(sc, plans)
    file_level = [i for i in result.issues if i.file == "열빠짐.xlsx"]
    assert len(file_level) == 1                          # 행마다 쏟아지지 않고 파일당 1건
    i = file_level[0]
    assert i.row is None and i.kind == "필수값 빈칸"
    assert "설비ID 열이 없어 이 파일은 중복 검사를 하지 못함" in i.message
    assert not any(i.kind == "중복 행" for i in result.issues)

    out = write_result(result, tmp_path / "out")
    es = load_workbook(out)["오류목록"]
    assert es.cell(2, 1).value == "열빠짐.xlsx" and es.cell(2, 2).value == "(전체)"


def test_two_candidates_for_one_standard_column(tmp_path):
    sc, _ = _stock_like(tmp_path)
    f = make_xlsx(tmp_path / "후보2개.xlsx", [
        ("Equip ID", "설비ID", "점검일", "온도", "상태", "장비번호"),
        ("X-1", "EQ-01", datetime(2026, 9, 1), 30, "정상", "Y-1"),
    ])
    plans = prepare(sc, [f])
    m = plans[0].match.get("설비ID")
    assert (m.source_index, m.method) == (1, "정확히 일치")   # 정확히 일치가 왼쪽 동의어보다 우선
    assert plans[0].match.unmatched_sources == [0, 5]
    warning = "[후보2개.xlsx] 설비ID 후보가 3개: B열 '설비ID' 사용, A열 'Equip ID' 무시, F열 '장비번호' 무시"
    assert pre_run_warnings(plans) == [warning]

    # 정확히 일치가 없으면 동의어 중 왼쪽 열
    g = make_xlsx(tmp_path / "동의어2개.xlsx", [
        ("장비번호", "점검일", "온도", "상태", "Equip ID"),
        ("EQ-01", datetime(2026, 9, 1), 30, "정상", "EQ-01"),
    ])
    assert pre_run_warnings(prepare(sc, [g])) == [
        "[동의어2개.xlsx] 설비ID 후보가 2개: A열 '장비번호' 사용, E열 'Equip ID' 무시"]

    # run_merge.py 출력에도 반드시 나온다
    code, out = run_merge_main([str(tmp_path / "s.yaml"), str(tmp_path), "--out", str(tmp_path / "out")])
    assert code == 0
    assert warning in out


def test_number_in_date_cell_suggestion_reaches_result_file(tmp_path):
    sc, _ = _stock_like(tmp_path)
    serial = (date(2026, 9, 2) - date(1899, 12, 30)).days
    f = make_xlsx(tmp_path / "숫자날짜.xlsx", [
        ("설비ID", "점검일", "온도", "상태"),
        ("EQ-01", 20260901, 30, "정상"),      # 8자리 숫자
        ("EQ-02", serial, 30, "정상"),        # 엑셀 날짜 일련번호가 숫자 서식으로 보이는 경우
        ("EQ-03", 7, 30, "정상"),             # 날짜로 볼 수 없는 숫자: 제안 없음
    ])
    result = execute(sc, prepare(sc, [f]))
    by_row = {i.row: i for i in result.issues}
    assert set(by_row) == {2, 3, 4} and all(i.kind == "형식 오류" for i in by_row.values())
    assert by_row[2].message.endswith("2026-09-01으로 수정 제안")
    assert by_row[3].message.endswith("2026-09-02으로 수정 제안")
    assert by_row[4].suggestion is None

    es = load_workbook(write_result(result, tmp_path / "out"))["오류목록"]
    header = [c.value for c in es[1]]
    sug_col = header.index("수정 제안값") + 1
    sug = {es.cell(r, 2).value: es.cell(r, sug_col).value for r in range(2, es.max_row + 1)
           if es.cell(r, 1).value}
    assert sug[2] == datetime(2026, 9, 1) and sug[3] == datetime(2026, 9, 2) and sug[4] is None
    # 취합결과에는 원래 숫자 그대로 (자동으로 고치지 않음)
    ws = load_workbook(write_result(result, tmp_path / "out"))["취합결과"]
    assert ws.cell(2, 4).value == 20260901

    # 수정 제안값 일괄 적용: 제안이 있는 셀만 바뀌고 초록색, 제안이 없는 셀은 원래 값과 오류 색
    wb = load_workbook(write_result(result, tmp_path / "out", apply_suggestions=True))
    ws = wb["취합결과"]
    assert ws.cell(2, 4).value == datetime(2026, 9, 1) and ws.cell(2, 4).fill.fgColor.rgb.endswith("C6EFCE")
    assert ws.cell(3, 4).value == datetime(2026, 9, 2)
    assert ws.cell(4, 4).value == 7 and ws.cell(4, 4).fill.fgColor.rgb.endswith("FFEB9C")
    es = wb["오류목록"]
    header = [c.value for c in es[1]]
    assert header[-1] == "처리"
    done = {es.cell(r, 2).value: es.cell(r, len(header)).value for r in range(2, es.max_row + 1)}
    assert done == {2: "자동 수정됨", 3: "자동 수정됨", 4: None}


def test_duplicate_group_numbers_in_error_sheet(tmp_path):
    sc, ok = _stock_like(tmp_path)
    dup = make_xlsx(tmp_path / "중복.xlsx", [
        ("설비ID", "점검일", "온도", "상태"),
        ("EQ-01", "2026-09-01", 31, "정상"),   # 정상.xlsx 2행과 중복 (날짜 표기만 다름)
        ("EQ-05", "2026-09-01", 31, "정상"),
        ("EQ-05", "2026.9.1", 32, "고장"),     # 같은 파일 안 중복
    ])
    result = execute(sc, prepare(sc, [ok, dup]))
    groups = {(i.file, i.row): i.dup_group for i in result.issues if i.kind == "중복 행"}
    assert groups == {("정상.xlsx", 2): "중복-1", ("중복.xlsx", 2): "중복-1",
                      ("중복.xlsx", 3): "중복-2", ("중복.xlsx", 4): "중복-2"}
    msg = {(i.file, i.row): i.message for i in result.issues if i.kind == "중복 행"}
    assert msg[("정상.xlsx", 2)].startswith("중복.xlsx 2행과")

    es = load_workbook(write_result(result, tmp_path / "out"))["오류목록"]
    header = [c.value for c in es[1]]
    g = header.index("중복 그룹") + 1
    sheet_groups = sorted(es.cell(r, g).value for r in range(2, es.max_row + 1) if es.cell(r, g).value)
    assert sheet_groups == ["중복-1", "중복-1", "중복-2", "중복-2"]
