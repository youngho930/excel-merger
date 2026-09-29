"""화면: 자유 양식 (직접 열 정하기). 체험 버튼 -> 머리글 확인 -> 열 묶기 -> 열 고르기 -> 실행 -> 결과·YAML."""

import json
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from streamlit.testing.v1 import AppTest

from answer_key import ROOT, compare, load_free_answer
from engine import ai_match
from engine.free_form import regroup
from engine.scenario import parse_scenario

import app

APP = str(ROOT / "app.py")
ANSWER = load_free_answer()
PAIRS = [("고객명", "성명"), ("연락처", "전화번호"), ("제품명", "모델"), ("증상", "고장내용"), ("접수일", "등록일")]


def free_demo() -> AppTest:
    at = AppTest.from_file(APP, default_timeout=120)
    at.run()
    at.button(key="free_demo_btn").click().run()
    assert not at.exception, at.exception
    return at


def groups_of(at):
    return regroup(at.session_state["ff_tables"], at.session_state["ff_merges"])


def gid(at, label):
    return next(g.gid for g in groups_of(at) if g.label == label)


def k(at, kind, label):
    return f"ff_{kind}_{at.session_state['load_id']}_{gid(at, label)}"


def merge(at, a, b):
    at.multiselect(key=f"ff_merge_pick_{at.session_state['load_id']}").set_value([gid(at, a), gid(at, b)])
    at.button(key="ff_merge_btn").click().run()
    assert not at.exception, at.exception


def configured() -> AppTest:
    """정답지 설정대로: 5쌍 묶기, 접수일을 맨 위로, 필수·중복기준 체크."""
    at = free_demo()
    for a, b in PAIRS:
        merge(at, a, b)
    at.button(key=k(at, "up", "접수일")).click().run()
    for label in ("접수일", "연락처", "제품명"):
        at.checkbox(key=k(at, "dup", label)).check().run()
    at.checkbox(key=k(at, "req", "고객명")).check().run()
    assert not at.exception, at.exception
    return at


def all_text(at) -> str:
    parts = []
    for kind in ("markdown", "caption", "warning", "error", "info", "success"):
        parts += [e.value for e in getattr(at, kind)]
    parts += [c.label for c in at.checkbox] + [b.label for b in at.button]
    return "\n".join(parts)


def test_first_screen_keeps_default_and_adds_free_button():
    at = AppTest.from_file(APP, default_timeout=120)
    at.run()
    sb = at.selectbox(key="scenario")
    assert sb.options[0] == app.FREE_LABEL                     # 목록 맨 위
    assert sb.value == "incoming_inspection"                   # 처음 선택은 그대로 첫 시나리오
    keys = [b.key for b in at.button]
    assert keys[:3] == ["sample_btn", "ai_demo_btn_0", "free_demo_btn"]
    assert at.button(key="free_demo_btn").label == "자유 양식 체험"


def test_free_demo_shows_headers_and_auto_groups():
    at = free_demo()
    assert at.selectbox(key="scenario").value == app.FREE_OPTION
    text = all_text(at)
    assert at.session_state["load_source"] == "free_demo"
    for f in ANSWER["files"]:
        assert f"**{app.md(f['file'])}** — 감지된 머리글: **{f['header_row']}행**" in text
    assert "확인 필요" not in text                             # 세 파일 모두 확신 높음
    assert len([n for n in at.number_input if n.key.startswith("ff_hr_")]) == 3
    assert len(groups_of(at)) == 13
    # 기본: 모든 묶음이 결과에 들어가고, 날짜만 있는 열은 '날짜로 비교'
    assert len(at.session_state["ff_order"]) == 13
    assert "날짜로 비교" in text
    # 수정 제안값·시나리오 표는 자유 양식에 없다
    assert "시나리오 자세히 보기" not in [e.label for e in at.expander]


def test_full_flow_matches_answer_key(tmp_path):
    at = configured()
    order = [g for g in at.session_state["ff_order"]]
    labels = {g.gid: g.label for g in groups_of(at)}
    assert [labels[g] for g in order] == ["접수일", "접수번호", "고객명", "연락처", "제품명", "증상", "담당기사", "주문번호"]
    assert at.checkbox(key=k(at, "req", "연락처")).proto.disabled      # 중복기준이면 필수가 켜지고 잠김
    assert at.checkbox(key=k(at, "req", "연락처")).value
    at.text_input(key=k(at, "name", "증상")).set_value("증상 내용").run()
    at.button(key="run_btn").click().run()
    assert not at.exception, at.exception
    result = at.session_state["result"]
    cmp = compare(result.issues, ANSWER["errors"])
    assert cmp.ok, cmp.report()
    c = result.counts()
    assert (c["필수값 빈칸"], c["중복 행"], sum(c.values())) == (5, 6, 11)
    # 결과 화면: 수정 제안값 항목은 숨기고, 중복 행 고르기는 그대로
    assert "apply_suggestions" not in [x.key for x in at.checkbox]
    assert "수정 제안 있음" not in [m.label for m in at.metric]
    assert len([x for x in at.checkbox if x.key and x.key.startswith("keep_")]) == 6
    # 결과 엑셀: 기존과 같은 시트 구조, 결과 열 이름·순서 반영
    wb = load_workbook(at.session_state["current_output"], read_only=True)
    assert wb.sheetnames == ["취합결과", "오류목록", "요약", "범례"]
    header = next(wb["취합결과"].iter_rows(max_row=1, values_only=True))
    assert list(header) == ["출처 파일", "원래 행", "접수일", "접수번호", "고객명", "연락처", "제품명", "증상 내용",
                            "담당기사", "주문번호"]
    wb.close()


def test_saved_yaml_preview_is_a_real_scenario(tmp_path):
    at = configured()
    at.text_input(key="ff_form_name").set_value("AS 접수 내역").run()
    at.text_input(key="ff_file_key").set_value("as_intake").run()
    code = next(c.value for c in at.code if c.language == "yaml")
    sc = parse_scenario(yaml.safe_load(code), key="as_intake")
    assert sc.name == "AS 접수 내역" and sc.dup_keys == ["접수일", "연락처", "제품명"]
    assert set(sc.column("고객명").aliases) == {"성명", "고객 명"}
    # 내려받기 전에 '형식: 날짜' 열을 보여준다
    assert any("'형식: 날짜'로 저장되는 열: 접수일" in i.value for i in at.info)
    assert at.get("download_button")
    # 기존 시나리오 key 와 같으면 경고
    at.text_input(key="ff_file_key").set_value("stock_count").run()
    assert any("기존 파일을 덮어씁니다" in w.value for w in at.warning)
    at.text_input(key="ff_file_key").set_value("한글이름").run()
    assert any("영문·숫자" in e.value for e in at.error)


@pytest.mark.parametrize("name, expect", [("원래 행", "'원래 행' 열과 겹쳐"), ("", "비어 있습니다"),
                                          ("접수번호", "두 번"), ("담당 기사2", None)])
def test_result_name_problems_block_run(name, expect):
    at = free_demo()
    at.text_input(key=k(at, "name", "고객명")).set_value(name).run()
    errors = " ".join(e.value for e in at.error)
    if expect is None:
        assert not errors and not at.button(key="run_btn").disabled
    else:
        assert app.md(expect) in errors or expect in errors, errors
        assert at.button(key="run_btn").disabled


def test_name_of_another_group_is_blocked():
    # 묶지 않은 '성명' 묶음이 있는데 고객명 묶음의 결과 열 이름을 '성명'으로 하면 그 열의 데이터를 잘못 가져온다
    at = free_demo()
    at.text_input(key=k(at, "name", "고객명")).set_value("성명").run()
    assert "다른 묶음" in " ".join(e.value for e in at.error)
    assert at.button(key="run_btn").disabled


def test_required_column_in_one_file_warns():
    at = free_demo()
    at.checkbox(key=k(at, "req", "담당기사")).check().run()
    warn = " ".join(w.value for w in at.warning)
    assert "파일 3개 중 1개에만" in warn and "나머지 2개 파일" in warn


def test_changing_header_row_rereads_file():
    at = free_demo()
    names = [t.file_name for t in at.session_state["ff_tables"]]
    fi = names.index("AS접수_콜센터.xlsx")
    key = f"ff_hr_{at.session_state['load_id']}_{fi}"
    at.number_input(key=key).set_value(1).run()
    t = at.session_state["ff_tables"][fi]
    assert (t.header_row, t.header_confidence) == (1, "지정")
    assert "직접 지정" in all_text(at)
    at.number_input(key=key).set_value(3).run()
    assert at.session_state["ff_tables"][fi].headers[:2] == ["등록일", "성명"]
    assert not at.exception


def test_ai_group_proposals_need_confirmation(monkeypatch):
    monkeypatch.setenv("GEMINI_API_KEY", "test-key")
    sent = []

    def fake(prompt, api_key, model, timeout):
        data = json.loads(prompt.rsplit("<data>", 1)[1].split("</data>")[0])
        sent.append(data)
        ids = {c["names"][0]: c["id"] for c in data["columns"]}
        return json.dumps({"groups": [[ids["고객명"], ids["성명"]], [ids["연락처"], ids["전화번호"]]]})
    monkeypatch.setattr(ai_match, "call_gemini", fake)
    at = free_demo()
    at.button(key="ff_ai_btn").click().run()
    assert not at.exception, at.exception
    assert at.session_state["ai_calls"] == 1
    # 열 이름만 보냈다 (데이터 값·파일 이름 없음)
    dumped = json.dumps(sent[0], ensure_ascii=False)
    assert "010-0000" not in dumped and ".xlsx" not in dumped
    # 아직 묶이지 않았다: 확인 후 적용
    assert len(groups_of(at)) == 13
    props = [c for c in at.checkbox if c.key and c.key.startswith("ff_prop_")]
    assert len(props) == 2 and all(c.value for c in props)
    props[1].uncheck().run()                                   # 두 번째는 따로 두기
    at.button(key="ff_apply_props").click().run()
    groups = {g.label: g for g in groups_of(at)}
    assert groups["고객명"].origin == "AI 추천" and "성명" in groups["고객명"].names
    assert "전화번호" in groups                                 # 따로 두기
    assert frozenset(("연락처", "전화번호")) in at.session_state["ff_rejected"]
    # 풀 수 있다
    at.selectbox(key=f"ff_split_pick_{at.session_state['load_id']}").set_value(groups["고객명"].gid).run()
    at.button(key="ff_split_btn").click().run()
    assert "성명" in {g.label for g in groups_of(at)}


HARNESS = '''
import sys
sys.path.insert(0, {root!r})
import streamlit as st
import app
app.init_state()
st.session_state.scenario = app.FREE_OPTION
if not st.session_state.get("tried"):
    st.session_state.tried = True
    app.load_files(None, [(n, open(p, "rb").read()) for n, p in {items!r}], "upload")
app.show_messages()
if st.session_state.ff_tables:
    app.render_free_steps(set())
'''
EVIL = "![p](https://attacker.invalid/b.png) [세션 만료](https://attacker.invalid/x)"


def test_user_text_is_escaped_in_free_form(tmp_path):
    rows = [(EVIL, "코드", "일자"), ("a", "x1", "2026.9.1"), ("b", "x2", "2026.13.02"), ("c", "x3", "2026.9.3")]
    wb = Workbook()
    for r in rows:
        wb.active.append(list(r))
    p = tmp_path / "[evil](x).xlsx"
    wb.save(p)
    at = AppTest.from_string(HARNESS.format(root=str(ROOT), items=[(p.name, str(p))]), default_timeout=120)
    at.run()
    assert not at.exception, at.exception
    code_key = next(t.key for t in at.text_input if t.key.startswith("ff_name_") and t.value == "코드")
    at.text_input(key=code_key).set_value("![q](https://attacker.invalid/q.png)").run()
    text = all_text(at)
    assert "attacker" in text                                   # 화면에 보이기는 하지만
    for raw in ("![p](", "[세션 만료](", "![q](", "[evil](x)"):
        assert raw not in text, raw                             # 마크다운으로 해석되지 않게 이스케이프
    # 날짜 안내의 예시 값도 이스케이프된 채로 보인다
    assert any(app.md("2026.13.02") in i.value for i in at.info)
    # st.html 에는 사용자 글자가 없다
    assert not [e for e in at.get("html") if "attacker" in e.proto.body]
