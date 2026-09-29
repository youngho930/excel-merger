"""화면(app.py) 테스트 — streamlit.testing.v1.AppTest.

샘플 체험 -> 실행 -> 결과 요약 건수, 수정 제안값 일괄 적용 on/off, 열 매칭 수정,
중복 그룹에서 남길 행 고르기, "남은 오류만 보기" 필터.
파일 읽기·실행은 실제 하위 프로세스(engine.jobs)로 돈다.
"""

from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

from answer_key import ROOT, load_answer
from engine.merge import AUTO_FIXED, DUP_RESOLVED, EXCLUDED

APP = str(ROOT / "app.py")
EXPECTED_TOTAL = {"incoming_inspection": 16, "stock_count": 15, "monthly_report": 11}
FIXED_RGB, BLUE_RGB = "C6EFCE", "BDD7EE"


def start(key: str) -> AppTest:
    at = AppTest.from_file(APP, default_timeout=120)
    at.run()
    assert not at.exception
    at.selectbox(key="scenario").set_value(key).run()
    at.button(key="sample_btn").click().run()
    ok(at)
    return at


def ok(at: AppTest) -> AppTest:
    assert not at.exception, at.exception
    assert not at.error, [e.value for e in at.error]
    return at


def run_merge(at: AppTest) -> AppTest:
    at.button(key="run_btn").click().run()
    return ok(at)


def metric(at: AppTest, label: str) -> str:
    return next(m.value for m in at.metric if m.label == label)


def num(at: AppTest, label: str) -> int:
    return int(metric(at, label).rstrip("건개행").replace(",", ""))


def kind_table(at: AppTest) -> pd.DataFrame:
    df = next(d.value for d in at.dataframe if list(d.value.columns) == ["오류 종류", "전체", "처리됨", "남은 오류"])
    return df.set_index("오류 종류")


def error_table(at: AppTest) -> pd.DataFrame:
    return at.dataframe[-1].value


def output(at: AppTest) -> Path:
    return Path(at.session_state["current_output"])


def keep_boxes(at: AppTest, group_index: int):
    prefix = f"keep_{at.session_state['result_id']}_{group_index}_"
    return [c for c in at.checkbox if c.key and c.key.startswith(prefix)]


def test_first_screen_has_no_errors():
    at = AppTest.from_file(APP, default_timeout=60)
    at.run()
    assert not at.exception and not at.error
    assert at.title[0].value == "엑셀 자동 취합·검증기"
    # 시나리오 3개 + 맨 위 "자유 양식 (직접 열 정하기)"
    options = at.selectbox(key="scenario").options
    assert len(options) == 4 and options[0] == "자유 양식 (직접 열 정하기)"
    assert not at.metric   # 파일을 올리기 전에는 결과가 없다


@pytest.mark.parametrize("key", sorted(EXPECTED_TOTAL))
def test_sample_run_summary(key):
    at = start(key)
    info = load_answer()["scenarios"][key]
    assert at.success[0].value == f"샘플 파일 {len(info['files'])}개를 읽었습니다."
    assert not at.warning   # 샘플은 매칭 경고가 없다
    run_merge(at)

    total = EXPECTED_TOTAL[key]
    assert metric(at, "파일 수") == f"{len(info['files'])}개"
    assert metric(at, "취합 행 수") == f"{sum(f['data_rows'] for f in info['files'])}행"
    assert metric(at, "전체 오류") == f"{total}건"
    assert metric(at, "처리됨") == "0건" and metric(at, "남은 오류") == f"{total}건"
    kt = kind_table(at)
    assert kt["전체"].sum() == total and kt["처리됨"].sum() == 0 and kt["남은 오류"].sum() == total

    # 오류 표: 기본은 "남은 오류만 보기" 켜짐 + 종류별 거르기
    assert at.checkbox(key="only_remaining").value is True
    df = error_table(at)
    assert len(df) == total and (df["처리"] == "").all()
    at.multiselect(key="kind_filter").set_value(["중복 행"]).run()
    df = error_table(at)
    assert len(df) == kt.loc["중복 행", "전체"] and set(df["오류 종류"]) == {"중복 행"}

    # 다운로드 버튼과 결과 파일 (제외한 행이 없으면 "제외된 행" 시트 없음, 처리 열은 빈칸)
    assert at.get("download_button")
    path = output(at)
    assert path.is_file() and path.parent.parent == Path(at.session_state["workspace"])
    assert load_workbook(path).sheetnames == ["취합결과", "오류목록", "요약", "범례"]
    errors = pd.read_excel(path, sheet_name="오류목록")
    assert len(errors) == total and errors["처리"].isna().all()
    assert not [c for c in errors.columns if str(c).startswith("Unnamed")]


def _cell(ws, file, row, column):
    header = [c.value for c in ws[1]]
    col = header.index(column) + 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == file and ws.cell(r, 2).value == row:
            return ws.cell(r, col)
    raise KeyError((file, row))


def _summary_kinds(wb) -> dict:
    rows = [[c.value for c in r] for r in wb["요약"].iter_rows()]
    head = next(i for i, r in enumerate(rows) if r[0] == "오류 종류")
    return {r[0]: tuple(r[1:4]) for r in rows[head + 1:] if r and r[0] and isinstance(r[1], int)}


def test_apply_suggestions_off_then_on():
    at = run_merge(start("incoming_inspection"))
    assert at.checkbox(key="apply_suggestions").value is False   # 기본값: 꺼짐
    assert metric(at, "수정 제안 있음") == "1건"

    # 꺼짐: 원래 값 그대로, 처리 열은 전부 빈칸
    off_path = output(at)
    off = load_workbook(off_path)
    cell = _cell(off["취합결과"], "수입검사_대성정공.xlsx", 14, "판정")
    assert cell.value == "합격 " and not cell.fill.fgColor.rgb.endswith(FIXED_RGB)
    assert pd.read_excel(off_path, sheet_name="오류목록")["처리"].isna().all()

    # 켬: 제안값으로 바꾸고 초록색, 오류목록 "처리" 열에 "자동 수정됨"
    at.checkbox(key="apply_suggestions").check().run()
    ok(at)
    on_path = output(at)
    assert on_path != off_path
    on = load_workbook(on_path)
    cell = _cell(on["취합결과"], "수입검사_대성정공.xlsx", 14, "판정")
    assert cell.value == "합격" and cell.fill.fgColor.rgb.endswith(FIXED_RGB)
    errors = pd.read_excel(on_path, sheet_name="오류목록")
    fixed = errors[errors["처리"] == AUTO_FIXED]
    assert list(zip(fixed["파일"], fixed["행"], fixed["기준열"])) == [("수입검사_대성정공.xlsx", 14, "판정")]
    assert fixed["값"].iloc[0] == "합격 " and fixed["오류 종류"].iloc[0] == "허용값 아닌 값"
    assert len(errors) == 16   # 전체 오류 건수는 그대로 (고친 사실만 남긴다)
    assert metric(at, "전체 오류") == "16건" and metric(at, "처리됨") == "1건" and metric(at, "남은 오류") == "15건"
    assert _summary_kinds(on)["합계"] == (16, 1, 15)

    # 남은 오류만 보기(기본)에서는 처리한 오류가 빠지고, 끄면 보인다
    assert len(error_table(at)) == 15
    at.checkbox(key="only_remaining").uncheck().run()
    df = error_table(at)
    assert len(df) == 16 and (df["처리"] == AUTO_FIXED).sum() == 1

    # 제안값이 없는 다른 오류 셀은 원래 값 그대로
    assert _cell(on["취합결과"], "수입검사_미래부품.xlsx", 7, "판정").value == "OK"

    # 다시 끄면 원래 값 파일
    at.checkbox(key="apply_suggestions").uncheck().run()
    assert at.checkbox(key="apply_suggestions").value is False
    assert output(at) == off_path


def test_exclude_one_row_from_duplicate_group():
    at = run_merge(start("stock_count"))
    before_rows = num(at, "취합 행 수")
    groups = at.session_state["result"].dup_groups()
    first_group = list(groups)[0]
    kept_row, dropped_row = groups[first_group]
    boxes = keep_boxes(at, 0)
    assert len(boxes) == 2 and all(b.value for b in boxes)   # 기본: 전부 남김

    boxes[1].uncheck().run()
    ok(at)
    assert num(at, "취합 행 수") == before_rows - 1
    assert metric(at, "제외한 행") == "1개"
    assert metric(at, "전체 오류") == "15건"            # 전체 건수는 그대로
    assert metric(at, "처리됨") == "2건" and metric(at, "남은 오류") == "13건"
    assert len(error_table(at)) == 13                   # 남은 오류만 보기
    assert not at.warning

    path = output(at)
    wb = load_workbook(path)
    assert wb.sheetnames == ["취합결과", "오류목록", "제외된 행", "요약", "범례"]
    ws = wb["취합결과"]
    assert ws.max_row - 1 == before_rows - 1
    rows = {(ws.cell(r, 1).value, ws.cell(r, 2).value) for r in range(2, ws.max_row + 1)}
    assert (dropped_row.source_file, dropped_row.excel_row) not in rows
    kept_cell = _cell(ws, kept_row.source_file, kept_row.excel_row, "출처 파일")
    assert not kept_cell.fill.fgColor.rgb.endswith(BLUE_RGB)   # 중복이 풀려 파란 색 뺌

    xs = wb["제외된 행"]
    rec = dict(zip([c.value for c in xs[1]], [c.value for c in xs[2]]))
    assert (rec["출처 파일"], rec["원래 행"], rec["중복 그룹"]) == \
           (dropped_row.source_file, dropped_row.excel_row, first_group)
    assert "중복 행" in rec["오류"]

    errors = pd.read_excel(path, sheet_name="오류목록")
    act = {(f, r): a for f, r, a in zip(errors["파일"], errors["행"], errors["처리"]) if isinstance(a, str)}
    assert act == {(dropped_row.source_file, dropped_row.excel_row): EXCLUDED,
                   (kept_row.source_file, kept_row.excel_row): DUP_RESOLVED}
    assert _summary_kinds(wb)["중복 행"] == (4, 2, 2)

    # 다시 남기면 A 단계와 같은 파일 구성
    keep_boxes(at, 0)[1].check().run()
    assert load_workbook(output(at)).sheetnames == ["취합결과", "오류목록", "요약", "범례"]


def test_excluding_whole_group_warns():
    at = run_merge(start("stock_count"))
    before_rows = num(at, "취합 행 수")
    for i in range(len(keep_boxes(at, 0))):
        keep_boxes(at, 0)[i].uncheck().run()
    ok(at)
    assert any("이 그룹의 행이 모두 빠집니다" in w.value for w in at.warning)
    assert num(at, "취합 행 수") == before_rows - 2
    assert load_workbook(output(at))["제외된 행"].max_row == 3


def test_apply_and_exclude_together():
    at = run_merge(start("incoming_inspection"))
    at.checkbox(key="apply_suggestions").check().run()
    keep_boxes(at, 0)[0].uncheck().run()
    ok(at)
    errors = pd.read_excel(output(at), sheet_name="오류목록")
    counts = errors["처리"].value_counts().to_dict()
    assert counts == {AUTO_FIXED: 1, EXCLUDED: 1, DUP_RESOLVED: 1}
    assert metric(at, "처리됨") == "3건" and metric(at, "남은 오류") == "13건"
    wb = load_workbook(output(at))
    assert _cell(wb["취합결과"], "수입검사_대성정공.xlsx", 14, "판정").value == "합격"
    assert _summary_kinds(wb)["합계"] == (16, 3, 13)


def test_matching_can_be_changed_on_screen():
    at = start("stock_count")
    load_id = at.session_state["load_id"]
    plans = at.session_state["plans"]
    fi = 0
    m = plans[fi].match
    si = next(i for i, cm in enumerate(m.matches) if cm.required)
    std = m.matches[si].standard
    first = m.matches[si].source_index
    other = next(cm.source_index for cm in m.matches if cm.standard != std and cm.source_index is not None)
    key = f"map_{load_id}_{fi}_{si}"
    assert at.selectbox(key=key).value == first

    # 필수 열의 매칭을 해제하면 누락 경고가 뜬다
    at.selectbox(key=key).set_value(None).run()
    assert not at.exception
    assert any(f"필수 열 '{std}'" in w.value for w in at.warning)
    assert at.session_state["plans"][fi].match.get(std).method == "매칭 안 됨"

    # 이미 다른 기준열에 쓰인 원본 열을 고르면 한국어 안내 후 되돌린다
    at.selectbox(key=key).set_value(other).run()
    assert any("이미" in e.value for e in at.error)
    assert at.selectbox(key=key).value is None

    # 원래 열로 되돌리면 "사용자 지정", 경고 사라짐 -> 실행 결과는 정답과 같다
    at.selectbox(key=key).set_value(first).run()
    assert at.session_state["plans"][fi].match.get(std).method == "사용자 지정"
    assert not at.warning
    run_merge(at)
    assert metric(at, "전체 오류") == "15건"


def test_changing_scenario_resets_screen():
    at = run_merge(start("monthly_report"))
    assert at.metric
    at.selectbox(key="scenario").set_value("stock_count").run()
    assert not at.exception
    assert not at.metric and at.session_state["plans"] is None


def test_upload_file_name_is_reduced_to_base_name():
    from app import safe_name
    assert safe_name("../../etc/passwd.xlsx") == "passwd.xlsx"
    assert safe_name("..\\..\\Windows\\a.xlsx") == "a.xlsx"
    assert safe_name("C:\\Users\\x\\b.xlsx") == "b.xlsx"
    assert safe_name("보고서.xlsx") == "보고서.xlsx"
    assert safe_name("..") == ".."   # 이런 이름은 load_files 가 거부한다


def test_upload_limit_matches_reader_limit():
    import tomllib
    from engine.reader import MAX_FILE_BYTES
    cfg = tomllib.loads((ROOT / ".streamlit" / "config.toml").read_text(encoding="utf-8"))
    assert cfg["server"]["maxUploadSize"] * 1024 * 1024 == MAX_FILE_BYTES
    assert cfg["client"]["showErrorDetails"] == "none"
