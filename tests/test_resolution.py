"""중복 행 제외 · 처리 열 · 요약 시트 (B 단계) — 엔진 단위 테스트.

확정된 규칙
- 뺀 행의 오류(다른 오류 + 그 행의 중복 오류)는 오류목록에 남기고 처리 열에 "행 제외됨"
- 처리 열은 항상 있다. 처리하지 않은 오류는 빈칸
- 제외 뒤 그룹에 1행만 남으면 그 행의 중복 오류는 "중복 해소(이 행을 남김)", 파란 색은 뺀다
- 2행 이상 남으면 그대로 (파란 색 유지, 처리 빈칸)
- 그룹 전체를 빼는 것은 허용 (fully_excluded_groups 로 알려 준다)
- 요약: 오류 종류별 전체 / 처리됨 / 남은 오류
"""

from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from engine import execute, load_scenario, prepare, write_result
from engine.merge import AUTO_FIXED, DUP_RESOLVED, EXCLUDED
from test_pipeline_cases import make_xlsx, write_yaml

BLUE, GREEN = "BDD7EE", "C6EFCE"
D = datetime(2026, 9, 1)
F = "A.xlsx"


@pytest.fixture
def result(tmp_path):
    sc = load_scenario(write_yaml(tmp_path / "eq.yaml"))
    path = make_xlsx(tmp_path / F, [
        ("설비ID", "점검일", "온도", "상태"),
        ("EQ-01", D, 30, "정상"),           # 2행  중복-1
        ("EQ-01", D, "28도", "정상 "),      # 3행  중복-1 + 형식 오류 + 허용값(제안 "정상")
        ("EQ-01", D, 31, "정상"),           # 4행  중복-1
        ("EQ-02", D, 30, "정상"),           # 5행  중복-2
        ("EQ-02", D, 30, "고장"),           # 6행  중복-2
        ("EQ-03", D, 30, "정상 "),          # 7행  허용값(제안 "정상")
    ])
    return execute(sc, prepare(sc, [path]))


def actions(result, res):
    """(행, 오류 종류) -> 처리."""
    return {(i.row, i.kind): a for i, a in zip(result.issues, res.actions)}


def result_rows(ws):
    return {ws.cell(r, 2).value: r for r in range(2, ws.max_row + 1)}


def summary(wb):
    """요약 시트: 오류 종류 -> (전체, 처리됨, 남은 오류), 그리고 위쪽 항목 -> 값."""
    ws = wb["요약"]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    head = next(i for i, r in enumerate(rows) if r[0] == "오류 종류")
    top = {r[0]: r[1] for r in rows[:head] if r and r[0]}
    kinds = {r[0]: tuple(r[1:4]) for r in rows[head + 1:] if r and r[0] and isinstance(r[1], int)}
    return top, kinds


def test_setup(result):
    assert result.counts() == {"필수값 빈칸": 0, "형식 오류": 1, "범위 밖 값": 0, "허용값 아닌 값": 2, "중복 행": 5}
    assert {g: [r.excel_row for r in rows] for g, rows in result.dup_groups().items()} == \
           {"중복-1": [2, 3, 4], "중복-2": [5, 6]}


def test_nothing_done_action_column_all_blank(result, tmp_path):
    res = result.resolve()
    assert set(res.actions) == {""} and res.totals() == (8, 0, 8)
    out = write_result(result, tmp_path)
    wb = load_workbook(out)
    assert wb.sheetnames == ["취합결과", "오류목록", "요약", "범례"]
    errors = pd.read_excel(out, sheet_name="오류목록")
    assert list(errors.columns)[-1] == "처리" and errors["처리"].isna().all()
    assert not [c for c in errors.columns if str(c).startswith("Unnamed")]
    assert wb["취합결과"].max_row - 1 == 6
    top, kinds = summary(wb)
    assert top["제외한 행 수"] == 0 and top["취합결과 행 수"] == 6
    assert kinds["중복 행"] == (5, 0, 5) and kinds["합계"] == (8, 0, 8)


def test_exclude_one_row_of_three(result, tmp_path):
    res = result.resolve(excluded=[(F, 3)])
    a = actions(result, res)
    # 뺀 행의 오류는 모두 "행 제외됨"
    assert a[(3, "형식 오류")] == a[(3, "허용값 아닌 값")] == a[(3, "중복 행")] == EXCLUDED
    # 2행 이상 남은 그룹은 그대로
    assert a[(2, "중복 행")] == a[(4, "중복 행")] == ""
    assert res.resolved_rows == frozenset() and res.fully_excluded_groups == []

    wb = load_workbook(write_result(result, tmp_path, excluded=[(F, 3)]))
    assert wb.sheetnames == ["취합결과", "오류목록", "제외된 행", "요약", "범례"]
    ws = wb["취합결과"]
    rows = result_rows(ws)
    assert ws.max_row - 1 == 5 and 3 not in rows
    assert ws.cell(rows[2], 1).fill.fgColor.rgb.endswith(BLUE)   # 아직 중복

    xs = wb["제외된 행"]
    header = [c.value for c in xs[1]]
    assert header[:4] == ["출처 파일", "원래 행", "중복 그룹", "오류"]
    rec = dict(zip(header, [c.value for c in xs[2]]))
    assert rec["출처 파일"] == F and rec["원래 행"] == 3 and rec["중복 그룹"] == "중복-1"
    assert rec["온도"] == "28도" and rec["상태"] == "정상 "          # 원래 값 그대로
    assert "형식 오류: 온도" in rec["오류"] and "중복 행" in rec["오류"]
    assert xs.max_row == 2

    top, kinds = summary(wb)
    assert top["제외한 행 수"] == 1 and top["취합결과 행 수"] == 5 and top["읽은 데이터 행 수"] == 6
    assert kinds["형식 오류"] == (1, 1, 0) and kinds["중복 행"] == (5, 1, 4) and kinds["합계"] == (8, 3, 5)


def test_group_left_with_one_row_is_resolved(result, tmp_path):
    ex = [(F, 3), (F, 4)]
    res = result.resolve(excluded=ex)
    a = actions(result, res)
    assert a[(2, "중복 행")] == DUP_RESOLVED
    assert a[(4, "중복 행")] == EXCLUDED
    assert a[(5, "중복 행")] == "" and a[(6, "중복 행")] == ""
    assert res.resolved_rows == frozenset({(F, 2)})

    ws = load_workbook(write_result(result, tmp_path, excluded=ex))["취합결과"]
    rows = result_rows(ws)
    header = [c.value for c in ws[1]]
    for c in (1, 2, header.index("설비ID") + 1, header.index("점검일") + 1):
        assert not ws.cell(rows[2], c).fill.fgColor.rgb.endswith(BLUE)   # 중복이 풀려 파란 색 뺌
    assert ws.cell(rows[5], 1).fill.fgColor.rgb.endswith(BLUE)          # 다른 그룹은 그대로


def test_whole_group_excluded_is_allowed(result, tmp_path):
    ex = [(F, 5), (F, 6)]
    res = result.resolve(excluded=ex)
    assert res.fully_excluded_groups == ["중복-2"]
    assert actions(result, res)[(5, "중복 행")] == EXCLUDED
    wb = load_workbook(write_result(result, tmp_path, excluded=ex))
    assert wb["취합결과"].max_row - 1 == 4 and wb["제외된 행"].max_row == 3


def test_apply_suggestions_together_with_exclusion(result, tmp_path):
    ex = [(F, 3)]
    res = result.resolve(excluded=ex, apply_suggestions=True)
    a = actions(result, res)
    assert a[(3, "허용값 아닌 값")] == EXCLUDED     # 뺀 행에는 제안값을 적용하지 않는다
    assert a[(7, "허용값 아닌 값")] == AUTO_FIXED
    assert (F, 3, "상태") not in res.fixes and res.fixes[(F, 7, "상태")] == "정상"
    assert res.totals() == (8, 4, 4)

    out = write_result(result, tmp_path, excluded=ex, apply_suggestions=True)
    wb = load_workbook(out)
    ws = wb["취합결과"]
    rows = result_rows(ws)
    col = [c.value for c in ws[1]].index("상태") + 1
    assert ws.cell(rows[7], col).value == "정상" and ws.cell(rows[7], col).fill.fgColor.rgb.endswith(GREEN)
    assert wb["제외된 행"].cell(2, 4 + 4).value == "정상 "   # 원래 값 (출처·행·그룹·오류 + 설비ID·점검일·온도·상태)
    errors = pd.read_excel(out, sheet_name="오류목록")
    assert errors["처리"].value_counts().to_dict() == {EXCLUDED: 3, AUTO_FIXED: 1}
    top, kinds = summary(wb)
    assert top["수정 제안값 일괄 적용"] == "예"
    assert kinds["허용값 아닌 값"] == (2, 2, 0) and kinds["합계"] == (8, 4, 4)


def test_resolve_does_not_change_result(result):
    before = [(i.row, i.kind, i.message) for i in result.issues]
    result.resolve(excluded=[(F, 2), (F, 3)], apply_suggestions=True)
    assert [(i.row, i.kind, i.message) for i in result.issues] == before
    assert len(result.rows) == 6
