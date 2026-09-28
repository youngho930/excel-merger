"""3개 시나리오 샘플을 엔진으로 돌려 정답지(expected_errors.json)와 대조한다."""

from collections import Counter
from pathlib import Path

import pytest
from openpyxl import load_workbook

from answer_key import ROOT, compare, load_answer
from engine import execute, load_scenario, prepare, write_result

ANSWER = load_answer()
SCENARIO_KEYS = sorted(ANSWER["scenarios"])


def run(key):
    sc = load_scenario(ROOT / "scenarios" / f"{key}.yaml")
    info = ANSWER["scenarios"][key]
    folder = ROOT / "samples" / info["folder"]
    # 파일 순서를 가나다순으로: 정답지(생성 순서)와 달라도 결과가 같아야 한다
    paths = sorted(folder / f["file"] for f in info["files"])
    plans = prepare(sc, paths)
    return sc, info, plans, execute(sc, plans)


def test_answer_key_has_37_errors():
    assert sum(len(s["errors"]) for s in ANSWER["scenarios"].values()) == 37


@pytest.mark.parametrize("key", SCENARIO_KEYS)
def test_engine_matches_answer_key(key):
    _, info, _, result = run(key)
    cmp = compare(result.issues, info["errors"])
    print(cmp.report())
    assert cmp.ok, "\n" + cmp.report()
    assert cmp.matched == len(info["errors"])


@pytest.mark.parametrize("key", SCENARIO_KEYS)
def test_no_issue_reported_twice(key):
    _, _, _, result = run(key)
    keys = Counter((i.file, i.row, i.standard, i.kind) for i in result.issues)
    assert [k for k, n in keys.items() if n > 1] == []


@pytest.mark.parametrize("key", SCENARIO_KEYS)
def test_header_rows_and_matching(key):
    _, info, plans, _ = run(key)
    by_name = {p.file_name: p for p in plans}
    for f in info["files"]:
        plan = by_name[f["file"]]
        assert plan.table.header_row == f["header_row"]
        assert plan.table.header_confidence == "높음"
        assert len(plan.table.rows) == f["data_rows"]
        assert plan.match.missing_required == []
        assert plan.match.missing_optional == []
        assert plan.match.unmatched_sources == []
        assert plan.match.conflicts == []
        assert plan.warnings() == []


@pytest.mark.parametrize("key", SCENARIO_KEYS)
def test_file_order_does_not_change_result(key):
    sc, info, plans, result = run(key)
    reversed_result = execute(sc, list(reversed(plans)))
    a = {(i.file, i.row, i.standard, i.kind) for i in result.issues}
    b = {(i.file, i.row, i.standard, i.kind) for i in reversed_result.issues}
    assert a == b


def test_duplicate_groups_pair_both_rows():
    _, _, _, result = run("stock_count")
    dups = [i for i in result.issues if i.kind == "중복 행"]
    groups = Counter(i.dup_group for i in dups)
    assert len(groups) == 2 and all(n == 2 for n in groups.values())
    pairs = {(i.file, i.row): i.message for i in dups}
    assert pairs[("재고실사_구미창고.xlsx", 13)].startswith("재고실사_평택1창고.xlsx 5행")
    assert pairs[("재고실사_평택1창고.xlsx", 5)].startswith("재고실사_구미창고.xlsx 13행")


def test_month_values_normalized_for_duplicates():
    # 날짜 셀 / "2026-08" / "2026년 8월"이 같은 달로 모여야 품질팀-생산팀 중복이 잡힌다
    _, _, _, result = run("monthly_report")
    dups = {(i.file, i.row) for i in result.issues if i.kind == "중복 행"}
    assert dups == {("월간실적_품질팀.xlsx", 12), ("월간실적_생산팀.xlsx", 12)}


def test_whitespace_suggestions():
    _, _, _, r1 = run("incoming_inspection")
    i = next(i for i in r1.issues if i.file == "수입검사_대성정공.xlsx" and i.row == 14)
    assert i.kind == "허용값 아닌 값" and i.value == "합격 "
    assert i.message == '앞뒤 공백 있음 → "합격"으로 수정 제안' and i.suggestion == "합격"
    _, _, _, r2 = run("monthly_report")
    i = next(i for i in r2.issues if i.file == "월간실적_영업팀.xlsx" and i.row == 10)
    assert i.message == '공백 차이 → "영업팀"일 가능성' and i.suggestion == "영업팀"


def test_nonexistent_dates_are_format_errors():
    _, _, _, r = run("incoming_inspection")
    i = next(i for i in r.issues if i.file == "수입검사_미래부품.xlsx" and i.row == 8)
    assert i.kind == "형식 오류" and "13월" in i.message
    _, _, _, r = run("stock_count")
    i = next(i for i in r.issues if i.file == "재고실사_구미창고.xlsx" and i.row == 6)
    assert i.kind == "형식 오류" and "9월 31일" in i.message


@pytest.mark.parametrize("key", SCENARIO_KEYS)
def test_result_workbook(key, tmp_path):
    sc, info, plans, result = run(key)
    sources = {p.table.path: p.table.path.read_bytes() for p in plans}
    out = write_result(result, tmp_path)
    assert out.parent == tmp_path.resolve()
    # 원본은 바이트 단위로 그대로
    for path, data in sources.items():
        assert path.read_bytes() == data

    wb = load_workbook(out)
    assert wb.sheetnames == ["취합결과", "오류목록", "범례"]
    ws = wb["취합결과"]
    header = [c.value for c in ws[1]]
    assert header == ["출처 파일", "원래 행"] + sc.column_names
    assert ws.max_row - 1 == sum(f["data_rows"] for f in info["files"])

    # 오류 셀(행 전체 제외)마다 색이 칠해져 있고, 값은 원래 값 그대로
    row_of = {(ws.cell(r, 1).value, ws.cell(r, 2).value): r for r in range(2, ws.max_row + 1)}
    col_of = {name: i for i, name in enumerate(header, start=1)}
    for issue in result.issues:
        if issue.row is None or issue.kind == "중복 행":
            continue
        cell = ws.cell(row_of[(issue.file, issue.row)], col_of[issue.standard])
        assert cell.fill.fgColor.rgb.endswith(("FFC7CE", "FFEB9C", "F8CBAD", "D9D2E9"))
        assert cell.value == issue.value

    es = wb["오류목록"]
    assert [c.value for c in es[1]][:9] == ["파일", "행", "열", "기준열", "오류 종류", "값", "설명",
                                            "수정 제안값", "중복 그룹"]
    assert sum(1 for r in range(2, es.max_row + 1) if es.cell(r, 1).value) == len(result.issues)
